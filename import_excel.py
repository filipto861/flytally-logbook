from __future__ import annotations

import argparse
import sqlite3
from datetime import datetime, date, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SCHEMA = """
PRAGMA foreign_keys = ON;

CREATE TABLE IF NOT EXISTS flights (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    evidence TEXT,
    registration TEXT,
    aircraft_type TEXT,
    aircraft_class TEXT,
    departure TEXT,
    arrival TEXT,
    off_block TEXT,
    takeoff TEXT,
    landing TEXT,
    on_block TEXT,
    starts INTEGER DEFAULT 1,
    commander TEXT,
    instructor TEXT,
    role TEXT,
    task TEXT,
    price_per_hour REAL,
    note TEXT
);

CREATE TABLE IF NOT EXISTS rates (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    registration TEXT NOT NULL,
    aircraft_type TEXT,
    valid_from TEXT,
    price_per_hour REAL,
    dry_price_per_hour REAL,
    source TEXT,
    UNIQUE(registration, valid_from)
);

CREATE TABLE IF NOT EXISTS flight_tracks (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    flight_id INTEGER NOT NULL,
    file_name TEXT,
    imported_at TEXT,
    point_count INTEGER,
    distance_km REAL,
    start_utc TEXT,
    end_utc TEXT,
    min_alt_m REAL,
    max_alt_m REAL,
    coordinates_json TEXT NOT NULL,
    FOREIGN KEY(flight_id) REFERENCES flights(id) ON DELETE CASCADE
);
"""

FLIGHT_HEADERS = {
    "Datum": "date",
    "Evidence": "evidence",
    "Imatrikulace": "registration",
    "Typ": "aircraft_type",
    "Třída": "aircraft_class",
    "Odlet": "departure",
    "Přílet": "arrival",
    "Off Block": "off_block",
    "Takeoff": "takeoff",
    "Landing": "landing",
    "On Block": "on_block",
    "Starty": "starts",
    "Velitel letadla": "commander",
    "Instruktor": "instructor",
    "Funkce": "role",
    "Úloha": "task",
    "Cena Kč/h": "price_per_hour",
    "Poznámka": "note",
}

RATE_HEADERS = {
    "Imatrikulace": "registration",
    "Typ": "aircraft_type",
    "Od data": "valid_from",
    "Cena Kč/h": "price_per_hour",
    "Suchá hodina Kč/h": "dry_price_per_hour",
    "Zdroj": "source",
}


def to_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text or None


def to_time(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.time().strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, timedelta):
        total_minutes = int(round(value.total_seconds() / 60)) % (24 * 60)
        return f"{total_minutes // 60:02d}:{total_minutes % 60:02d}"
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).strftime("%H:%M")
        except ValueError:
            pass
    return text


def normalize_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def normalize_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except Exception:
        text = str(value).replace(" ", "").replace("Kč/h", "").replace("Kč", "").replace(",", ".")
        try:
            return float(text)
        except Exception:
            return None


def normalize_int(value: Any) -> int:
    if value is None or value == "":
        return 0
    try:
        return int(value)
    except Exception:
        try:
            return int(float(str(value).replace(",", ".")))
        except Exception:
            return 0


def import_workbook(xlsx_path: str | Path, db_path: str | Path, reset: bool = True) -> None:
    xlsx_path = Path(xlsx_path)
    db_path = Path(db_path)
    db_path.parent.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(xlsx_path, data_only=True)
    con = sqlite3.connect(db_path)
    cur = con.cursor()
    cur.executescript(SCHEMA)
    if reset:
        cur.execute("DELETE FROM flight_tracks")
        cur.execute("DELETE FROM flights")
        cur.execute("DELETE FROM rates")
        cur.execute("DELETE FROM sqlite_sequence WHERE name IN ('flight_tracks','flights','rates')")

    if "Zápisník letů" in wb.sheetnames:
        ws = wb["Zápisník letů"]
        header_row = [cell.value for cell in ws[1]]
        col_map = {FLIGHT_HEADERS[h]: idx + 1 for idx, h in enumerate(header_row) if h in FLIGHT_HEADERS}
        fields = [
            "date", "evidence", "registration", "aircraft_type", "aircraft_class",
            "departure", "arrival", "off_block", "takeoff", "landing", "on_block",
            "starts", "commander", "instructor", "role", "task", "price_per_hour", "note",
        ]
        for row in range(2, ws.max_row + 1):
            raw_date = ws.cell(row, col_map.get("date", 1)).value
            flight_date = to_date(raw_date)
            if not flight_date:
                continue
            data = {}
            for field in fields:
                col = col_map.get(field)
                value = ws.cell(row, col).value if col else None
                if field == "date":
                    data[field] = flight_date
                elif field in {"off_block", "takeoff", "landing", "on_block"}:
                    data[field] = to_time(value)
                elif field == "starts":
                    data[field] = normalize_int(value)
                elif field == "price_per_hour":
                    data[field] = normalize_number(value)
                elif field in {"registration", "evidence", "role", "aircraft_class", "departure", "arrival"}:
                    txt = normalize_text(value)
                    data[field] = txt.upper() if txt else None
                else:
                    data[field] = normalize_text(value)
            cur.execute(
                """
                INSERT INTO flights (
                    date, evidence, registration, aircraft_type, aircraft_class,
                    departure, arrival, off_block, takeoff, landing, on_block,
                    starts, commander, instructor, role, task, price_per_hour, note
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [data[f] for f in fields],
            )

    if "Ceník" in wb.sheetnames:
        ws = wb["Ceník"]
        header_row = [cell.value for cell in ws[1]]
        col_map = {RATE_HEADERS[h]: idx + 1 for idx, h in enumerate(header_row) if h in RATE_HEADERS}
        fields = ["registration", "aircraft_type", "valid_from", "price_per_hour", "dry_price_per_hour", "source"]
        for row in range(2, ws.max_row + 1):
            registration_col = col_map.get("registration")
            registration = normalize_text(ws.cell(row, registration_col).value) if registration_col else None
            if not registration:
                continue
            data = {}
            for field in fields:
                col = col_map.get(field)
                value = ws.cell(row, col).value if col else None
                if field == "valid_from":
                    data[field] = to_date(value) or normalize_text(value)
                elif field in {"price_per_hour", "dry_price_per_hour"}:
                    data[field] = normalize_number(value)
                elif field == "registration":
                    data[field] = normalize_text(value).upper()
                else:
                    data[field] = normalize_text(value)
            cur.execute(
                """
                INSERT OR REPLACE INTO rates (
                    registration, aircraft_type, valid_from, price_per_hour, dry_price_per_hour, source
                ) VALUES (?, ?, ?, ?, ?, ?)
                """,
                [data[f] for f in fields],
            )

    con.commit()
    con.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Import Excel logbook into SQLite database.")
    parser.add_argument("xlsx", help="Path to Excel workbook")
    parser.add_argument("db", nargs="?", default="data/logbook.sqlite", help="SQLite database path")
    parser.add_argument("--append", action="store_true", help="Append instead of reset")
    args = parser.parse_args()
    import_workbook(args.xlsx, args.db, reset=not args.append)
