from __future__ import annotations

import html
import math
import re
from datetime import date, datetime
from io import BytesIO
from typing import Any

import pandas as pd
import streamlit as st

from .config import LOCAL_TZ
from .metrics import build_summary, fmt_minutes, fmt_money

def _export_date_bounds(df: pd.DataFrame) -> tuple[date, date]:
    if df.empty or "date_dt" not in df.columns or df["date_dt"].dropna().empty:
        today = date.today()
        return today, today
    vals = pd.to_datetime(df["date_dt"], errors="coerce").dropna()
    return vals.min().date(), vals.max().date()

def _safe_filename_part(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[^a-z0-9_\-]+", "-", text, flags=re.IGNORECASE).strip("-")
    return text or "export"

def _export_prefix(df: pd.DataFrame, label: str = "logbook") -> str:
    if df.empty or "date_dt" not in df.columns or df["date_dt"].dropna().empty:
        return _safe_filename_part(label)
    vals = pd.to_datetime(df["date_dt"], errors="coerce").dropna()
    start = vals.min().strftime("%Y%m%d")
    end = vals.max().strftime("%Y%m%d")
    return f"{_safe_filename_part(label)}_{start}_{end}"

def _duration_hours(minutes: Any) -> float:
    if minutes is None or pd.isna(minutes):
        return 0.0
    return round(float(minutes) / 60.0, 2)

def _money_number(value: Any) -> float:
    if value is None or pd.isna(value):
        return 0.0
    return round(float(value), 0)

@st.cache_data(show_spinner=False, ttl=300)
def make_logbook_export_df(df: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "Datum", "Evidence", "Imatrikulace", "Typ", "Třída", "Odlet", "Přílet",
        "Off Block", "Vzlet", "Přistání", "On Block", "Block Time", "Air Time",
        "Block h", "Air h", "Starty", "Velitel", "Instruktor", "Funkce", "Úloha",
        "Účtování", "Cena Kč/h", "Cena letu Kč", "GPS tracky", "GPS km", "Poznámka",
    ]
    if df.empty:
        return pd.DataFrame(columns=columns)
    work = df.sort_values(["date_dt", "off_block", "id"], na_position="last").copy()
    rows = []
    for _, r in work.iterrows():
        rows.append({
            "Datum": r.get("date"),
            "Evidence": r.get("evidence"),
            "Imatrikulace": r.get("registration"),
            "Typ": r.get("aircraft_type"),
            "Třída": r.get("aircraft_class"),
            "Odlet": r.get("departure"),
            "Přílet": r.get("arrival"),
            "Off Block": r.get("off_block"),
            "Vzlet": r.get("takeoff"),
            "Přistání": r.get("landing"),
            "On Block": r.get("on_block"),
            "Block Time": fmt_minutes(r.get("block_minutes")),
            "Air Time": fmt_minutes(r.get("air_minutes")),
            "Block h": _duration_hours(r.get("block_minutes")),
            "Air h": _duration_hours(r.get("air_minutes")),
            "Starty": int(r.get("starts") or 0),
            "Velitel": r.get("commander"),
            "Instruktor": r.get("instructor"),
            "Funkce": r.get("role"),
            "Úloha": r.get("task"),
            "Účtování": r.get("billing_basis"),
            "Cena Kč/h": _money_number(r.get("price_per_hour")),
            "Cena letu Kč": _money_number(r.get("cost")),
            "GPS tracky": int(r.get("track_count") or 0),
            "GPS km": round(float(r.get("gps_km") or 0), 1),
            "Poznámka": r.get("note"),
        })
    return pd.DataFrame(rows, columns=columns)

def _minutes_for_role(df: pd.DataFrame, role: str) -> int:
    if df.empty:
        return 0
    return int(df["block_minutes"].where(df["role"].eq(role), 0).fillna(0).sum())

@st.cache_data(show_spinner=False, ttl=300)
def make_summary_table(df: pd.DataFrame) -> pd.DataFrame:
    s = build_summary(df)
    rows = [
        ("Počet letů", s["flights"]),
        ("Starty", s["starts"]),
        ("Block Time", fmt_minutes(s["total"])),
        ("Air Time", fmt_minutes(s["air"])),
        ("PIC", fmt_minutes(s["pic"])),
        ("PIC ULL", fmt_minutes(s["pic_ull"])),
        ("PIC EASA", fmt_minutes(s["pic_easa"])),
        ("DUAL", fmt_minutes(s["dual"])),
        ("Safety Pilot", fmt_minutes(s["safety"])),
        ("ULL celkem", fmt_minutes(s["ull"])),
        ("EASA celkem", fmt_minutes(s["easa"])),
        ("GPS tracky", s["tracks"]),
        ("GPS km", round(float(s["gps_km"]), 1)),
        ("Náklady", fmt_money(float(s["cost"]))),
    ]
    return pd.DataFrame(rows, columns=["Metrika", "Hodnota"])

@st.cache_data(show_spinner=False, ttl=300)
def make_group_summary(df: pd.DataFrame, group_cols: list[str]) -> pd.DataFrame:
    base_cols = group_cols + ["Lety", "Starty", "Block", "Air", "PIC", "DUAL", "Safety", "GPS km", "Náklady Kč"]
    if df.empty:
        return pd.DataFrame(columns=base_cols)
    work = df.copy()
    for col in group_cols:
        if col not in work.columns:
            work[col] = ""
        work[col] = work[col].fillna("").astype(str).replace("", "—")
    work["_block"] = work["block_minutes"].fillna(0)
    work["_air"] = work["air_minutes"].fillna(0)
    work["_pic"] = work["_block"].where(work["role"].eq("PIC"), 0)
    work["_dual"] = work["_block"].where(work["role"].eq("DUAL"), 0)
    work["_safety"] = work["_block"].where(work["role"].eq("SAFETY PILOT"), 0)
    grouped = work.groupby(group_cols, dropna=False).agg(
        Lety=("id", "count"),
        Starty=("starts", "sum"),
        BlockMin=("_block", "sum"),
        AirMin=("_air", "sum"),
        PicMin=("_pic", "sum"),
        DualMin=("_dual", "sum"),
        SafetyMin=("_safety", "sum"),
        GpsKm=("gps_km", "sum"),
        Cost=("cost", "sum"),
    ).reset_index()
    grouped["Block"] = grouped["BlockMin"].apply(fmt_minutes)
    grouped["Air"] = grouped["AirMin"].apply(fmt_minutes)
    grouped["PIC"] = grouped["PicMin"].apply(fmt_minutes)
    grouped["DUAL"] = grouped["DualMin"].apply(fmt_minutes)
    grouped["Safety"] = grouped["SafetyMin"].apply(fmt_minutes)
    grouped["GPS km"] = grouped["GpsKm"].fillna(0).round(1)
    grouped["Náklady Kč"] = grouped["Cost"].fillna(0).round(0)
    grouped = grouped.sort_values(["BlockMin", "Lety"], ascending=False)
    return grouped[base_cols]

@st.cache_data(show_spinner=False, ttl=300)
def make_route_summary(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=["Trasa", "Lety", "Starty", "Block", "Air", "GPS km", "Náklady Kč"])
    work = df.copy()
    work["Trasa"] = work["departure"].fillna("").astype(str).str.upper().str.strip() + "–" + work["arrival"].fillna("").astype(str).str.upper().str.strip()
    work.loc[work["Trasa"].eq("–"), "Trasa"] = "—"
    grouped = work.groupby("Trasa", dropna=False).agg(
        Lety=("id", "count"),
        Starty=("starts", "sum"),
        BlockMin=("block_minutes", "sum"),
        AirMin=("air_minutes", "sum"),
        GpsKm=("gps_km", "sum"),
        Cost=("cost", "sum"),
    ).reset_index()
    grouped["Block"] = grouped["BlockMin"].apply(fmt_minutes)
    grouped["Air"] = grouped["AirMin"].apply(fmt_minutes)
    grouped["GPS km"] = grouped["GpsKm"].fillna(0).round(1)
    grouped["Náklady Kč"] = grouped["Cost"].fillna(0).round(0)
    grouped = grouped.sort_values(["Lety", "BlockMin"], ascending=False)
    return grouped[["Trasa", "Lety", "Starty", "Block", "Air", "GPS km", "Náklady Kč"]]

@st.cache_data(show_spinner=False, ttl=300)
def make_airport_summary(df: pd.DataFrame) -> pd.DataFrame:
    columns = ["Letiště", "Návštěvy", "Odlety", "Přílety", "První let", "Poslední let"]
    if df.empty:
        return pd.DataFrame(columns=columns)
    rows = []
    for kind, col in [("Odlety", "departure"), ("Přílety", "arrival")]:
        tmp = df[["date", col]].copy()
        tmp["Letiště"] = tmp[col].fillna("").astype(str).str.upper().str.strip()
        tmp = tmp[tmp["Letiště"].ne("")]
        tmp["Odlety"] = 1 if kind == "Odlety" else 0
        tmp["Přílety"] = 1 if kind == "Přílety" else 0
        rows.append(tmp[["date", "Letiště", "Odlety", "Přílety"]])
    if not rows:
        return pd.DataFrame(columns=columns)
    work = pd.concat(rows, ignore_index=True)
    grouped = work.groupby("Letiště", dropna=False).agg(
        Odlety=("Odlety", "sum"),
        Přílety=("Přílety", "sum"),
        První_let=("date", "min"),
        Poslední_let=("date", "max"),
    ).reset_index()
    grouped["Návštěvy"] = grouped["Odlety"] + grouped["Přílety"]
    grouped = grouped.sort_values(["Návštěvy", "Letiště"], ascending=[False, True])
    grouped = grouped.rename(columns={"První_let": "První let", "Poslední_let": "Poslední let"})
    return grouped[columns]

def _excel_safe_value(value: Any) -> Any:
    # OpenPyXL cannot reliably style/save NaN, ±Inf or pandas missing values.
    # Export should never fail because one optional numeric field is empty.
    try:
        if value is None or pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value

def _append_dataframe(ws, df: pd.DataFrame, start_row: int = 1) -> None:
    for col_idx, col in enumerate(df.columns, start=1):
        ws.cell(start_row, col_idx, col)
    for row_idx, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row_idx, col_idx, _excel_safe_value(value))

def _style_export_sheet(ws, title: str | None = None) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(color="0F172A", bold=True, size=14)
    thin = Side(style="thin", color="D1D5DB")
    if title:
        ws.insert_rows(1)
        ws.cell(1, 1, title)
        ws.cell(1, 1).font = title_font
        ws.row_dimensions[1].height = 22
        header_row = 2
    else:
        header_row = 1
    if ws.max_row >= header_row:
        for cell in ws[header_row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin)
    ws.freeze_panes = f"A{header_row + 1}"
    if ws.max_column and ws.max_row >= header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="E5E7EB"))
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                try:
                    numeric_value = float(cell.value)
                    if math.isfinite(numeric_value):
                        cell.number_format = '#,##0.00' if abs(numeric_value - int(numeric_value)) > 0.001 else '#,##0'
                except (TypeError, ValueError, OverflowError):
                    pass
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = max(len(str(ws.cell(row, col).value or "")) for row in range(1, min(ws.max_row, 300) + 1))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 34)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

@st.cache_data(show_spinner=False, ttl=300)
def export_excel(df: pd.DataFrame) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Zápisník"
    detail = make_logbook_export_df(df)
    _append_dataframe(ws, detail)
    _style_export_sheet(ws, "Letový zápisník")

    ws2 = wb.create_sheet("Souhrn")
    _append_dataframe(ws2, make_summary_table(df))
    _style_export_sheet(ws2, "Souhrn")

    sheets = [
        ("Letadla", make_group_summary(df, ["registration", "aircraft_type", "evidence"]), "Souhrn podle letadel"),
        ("Funkce", make_group_summary(df, ["role"]), "Souhrn podle funkce"),
        ("Trasy", make_route_summary(df), "Souhrn tras"),
        ("Letiště", make_airport_summary(df), "Souhrn letišť"),
    ]
    for name, table, title in sheets:
        wsx = wb.create_sheet(name)
        _append_dataframe(wsx, table)
        _style_export_sheet(wsx, title)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()

@st.cache_data(show_spinner=False, ttl=300)
def build_print_html(df: pd.DataFrame, title: str = "Letový zápisník") -> str:
    summary = build_summary(df)
    detail = make_logbook_export_df(df)
    generated = datetime.now(LOCAL_TZ).strftime("%d.%m.%Y %H:%M")
    if df.empty or "date_dt" not in df.columns or df["date_dt"].dropna().empty:
        period = "—"
    else:
        vals = pd.to_datetime(df["date_dt"], errors="coerce").dropna()
        period = f"{vals.min().strftime('%d.%m.%Y')} – {vals.max().strftime('%d.%m.%Y')}"

    cards = [
        ("Lety", summary["flights"]),
        ("Starty", summary["starts"]),
        ("Block", fmt_minutes(summary["total"])),
        ("Air", fmt_minutes(summary["air"])),
        ("PIC", fmt_minutes(summary["pic"])),
        ("DUAL", fmt_minutes(summary["dual"])),
        ("GPS km", round(float(summary["gps_km"]), 1)),
        ("Náklady", fmt_money(float(summary["cost"]))),
    ]
    cards_html = "".join(f"<div class='card'><span>{html.escape(str(label))}</span><strong>{html.escape(str(value))}</strong></div>" for label, value in cards)

    columns = ["Datum", "Imatrikulace", "Typ", "Odlet", "Přílet", "Off Block", "Vzlet", "Přistání", "On Block", "Block Time", "Air Time", "Starty", "Funkce", "Úloha"]
    header_html = "".join(f"<th>{html.escape(col)}</th>" for col in columns)
    row_html = []
    for _, r in detail.iterrows():
        row_html.append("<tr>" + "".join(f"<td>{html.escape(str(r.get(col) or ''))}</td>" for col in columns) + "</tr>")
    body_html = "".join(row_html) or f"<tr><td colspan='{len(columns)}'>Žádná data.</td></tr>"

    css = """
    body { font-family: Inter, Segoe UI, Arial, sans-serif; color:#111827; margin:28px; }
    h1 { margin:0 0 4px 0; font-size:24px; }
    .meta { color:#4b5563; font-size:12px; margin-bottom:18px; }
    .cards { display:grid; grid-template-columns:repeat(4, 1fr); gap:8px; margin:16px 0 18px; }
    .card { border:1px solid #d1d5db; border-radius:10px; padding:10px 12px; background:#f9fafb; }
    .card span { display:block; color:#6b7280; font-size:11px; text-transform:uppercase; letter-spacing:.04em; }
    .card strong { display:block; margin-top:4px; font-size:17px; }
    table { width:100%; border-collapse:collapse; font-size:10px; }
    th { background:#111827; color:white; padding:7px 6px; text-align:left; }
    td { border-bottom:1px solid #e5e7eb; padding:5px 6px; vertical-align:top; }
    tr:nth-child(even) td { background:#f9fafb; }
    @media print {
      body { margin:10mm; }
      .cards { grid-template-columns:repeat(4, 1fr); }
      table { font-size:8.5px; }
      th, td { padding:4px; }
    }
    """
    return f"""<!doctype html>
<html lang="cs">
<head><meta charset="utf-8"><title>{html.escape(title)}</title><style>{css}</style></head>
<body>
  <h1>{html.escape(title)}</h1>
  <div class="meta">Období: {html.escape(period)} · Vygenerováno: {html.escape(generated)}</div>
  <div class="cards">{cards_html}</div>
  <table><thead><tr>{header_html}</tr></thead><tbody>{body_html}</tbody></table>
</body>
</html>"""

