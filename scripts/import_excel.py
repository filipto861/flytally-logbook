from __future__ import annotations

import argparse
import sqlite3
from datetime import datetime, date, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from logbook_core.schema import SCHEMA
from logbook_core.tenancy import ensure_tenancy_schema
from logbook_core.permissions import strict_user_id
from logbook_core.performance import apply_sqlite_pragmas

FLIGHT_HEADERS = {
    "Datum": "date", "Evidence": "evidence", "Imatrikulace": "registration",
    "Typ": "aircraft_type", "Třída": "aircraft_class", "Odlet": "departure",
    "Přílet": "arrival", "Off Block": "off_block", "Takeoff": "takeoff",
    "Landing": "landing", "On Block": "on_block", "Starty": "starts",
    "Velitel letadla": "commander", "Instruktor": "instructor", "Funkce": "role",
    "Úloha": "task", "Cena Kč/h": "price_per_hour", "Poznámka": "note",
}
RATE_HEADERS = {
    "Imatrikulace": "registration", "Typ": "aircraft_type", "Od data": "valid_from",
    "Cena Kč/h": "price_per_hour", "Suchá hodina Kč/h": "dry_price_per_hour", "Zdroj": "source",
}


def to_date(value: Any) -> str | None:
    if value is None or value == "": return None
    if isinstance(value, datetime): return value.date().isoformat()
    if isinstance(value, date): return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y"):
        try: return datetime.strptime(text, fmt).date().isoformat()
        except ValueError: pass
    return text or None


def to_time(value: Any) -> str | None:
    if value is None or value == "": return None
    if isinstance(value, datetime): return value.time().strftime("%H:%M")
    if isinstance(value, time): return value.strftime("%H:%M")
    if isinstance(value, timedelta):
        total_minutes = int(round(value.total_seconds() / 60)) % (24 * 60)
        return f"{total_minutes // 60:02d}:{total_minutes % 60:02d}"
    text = str(value).strip()
    if not text: return None
    for fmt in ("%H:%M", "%H:%M:%S"):
        try: return datetime.strptime(text, fmt).strftime("%H:%M")
        except ValueError: pass
    return text


def normalize_text(value: Any) -> str | None:
    if value is None: return None
    text = str(value).strip()
    return text if text else None


def normalize_number(value: Any) -> float | None:
    if value is None or value == "": return None
    try: return float(value)
    except Exception:
        text = str(value).replace(" ", "").replace("Kč/h", "").replace("Kč", "").replace(",", ".")
        try: return float(text)
        except Exception: return None


def normalize_int(value: Any) -> int:
    if value is None or value == "": return 0
    try: return int(value)
    except Exception:
        try: return int(float(str(value).replace(",", ".")))
        except Exception: return 0


def import_workbook(
    xlsx_path: str | Path,
    db_path: str | Path,
    *,
    user_id: int = 1,
    reset: bool = True,
) -> None:
    """Import legacy Excel data into exactly one Logbook tenant.

    v0.69 deliberately removes the old global reset behavior. Even when
    ``reset=True``, only data owned by ``user_id`` is replaced.
    """
    uid = strict_user_id(user_id)
    xlsx_path = Path(xlsx_path)
    db_path = Path(db_path)
    db_path.parent.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(xlsx_path, data_only=True)
    con = sqlite3.connect(db_path, timeout=10.0)
    con.row_factory = sqlite3.Row
    apply_sqlite_pragmas(con, initial=True)
    con.executescript(SCHEMA)
    ensure_tenancy_schema(con)
    if not con.execute("SELECT 1 FROM users WHERE id = ?", (uid,)).fetchone():
        con.close()
        raise ValueError(f"User ID {uid} v databázi neexistuje.")

    cur = con.cursor()
    if reset:
        # Child tables first. Never touch another user's rows.
        cur.execute("DELETE FROM track_points WHERE user_id = ?", (uid,))
        cur.execute("DELETE FROM flight_tracks WHERE user_id = ?", (uid,))
        cur.execute("DELETE FROM flights WHERE user_id = ?", (uid,))
        cur.execute("DELETE FROM rates WHERE user_id = ?", (uid,))

    if "Zápisník letů" in wb.sheetnames:
        ws = wb["Zápisník letů"]
        header_row = [cell.value for cell in ws[1]]
        col_map = {FLIGHT_HEADERS[h]: idx + 1 for idx, h in enumerate(header_row) if h in FLIGHT_HEADERS}
        fields = [
            "date", "evidence", "registration", "aircraft_type", "aircraft_class",
            "departure", "arrival", "off_block", "takeoff", "landing", "on_block",
            "starts", "commander", "instructor", "role", "task", "price_per_hour", "note",
        ]
        for row in range(2, ws.max_row + 1):
            raw_date = ws.cell(row, col_map.get("date", 1)).value
            flight_date = to_date(raw_date)
            if not flight_date: continue
            data: dict[str, Any] = {}
            for field in fields:
                col = col_map.get(field)
                value = ws.cell(row, col).value if col else None
                if field == "date": data[field] = flight_date
                elif field in {"off_block", "takeoff", "landing", "on_block"}: data[field] = to_time(value)
                elif field == "starts": data[field] = normalize_int(value)
                elif field == "price_per_hour": data[field] = normalize_number(value)
                elif field in {"registration", "evidence", "role", "aircraft_class", "departure", "arrival"}:
                    txt = normalize_text(value); data[field] = txt.upper() if txt else None
                else: data[field] = normalize_text(value)
            cur.execute(
                """
                INSERT INTO flights (
                    user_id, date, evidence, registration, aircraft_type, aircraft_class,
                    departure, arrival, off_block, takeoff, landing, on_block,
                    starts, commander, instructor, role, task, price_per_hour, note
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [uid, *[data[f] for f in fields]],
            )

    if "Ceník" in wb.sheetnames:
        ws = wb["Ceník"]
        header_row = [cell.value for cell in ws[1]]
        col_map = {RATE_HEADERS[h]: idx + 1 for idx, h in enumerate(header_row) if h in RATE_HEADERS}
        fields = ["registration", "aircraft_type", "valid_from", "price_per_hour", "dry_price_per_hour", "source"]
        for row in range(2, ws.max_row + 1):
            registration_col = col_map.get("registration")
            registration = normalize_text(ws.cell(row, registration_col).value) if registration_col else None
            if not registration: continue
            data: dict[str, Any] = {}
            for field in fields:
                col = col_map.get(field)
                value = ws.cell(row, col).value if col else None
                if field == "valid_from": data[field] = to_date(value) or normalize_text(value)
                elif field in {"price_per_hour", "dry_price_per_hour"}: data[field] = normalize_number(value)
                elif field == "registration":
                    txt = normalize_text(value); data[field] = txt.upper() if txt else None
                else: data[field] = normalize_text(value)
            cur.execute(
                """
                INSERT INTO rates (
                    user_id, registration, aircraft_type, valid_from,
                    price_per_hour, dry_price_per_hour, source
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(user_id, registration, valid_from) DO UPDATE SET
                    aircraft_type=excluded.aircraft_type,
                    price_per_hour=excluded.price_per_hour,
                    dry_price_per_hour=excluded.dry_price_per_hour,
                    source=excluded.source
                """,
                [uid, *[data[f] for f in fields]],
            )

    con.commit()
    con.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Import Excel logbook into one Logbook user profile.")
    parser.add_argument("xlsx", help="Path to Excel workbook")
    parser.add_argument("db", nargs="?", default="data/logbook.sqlite", help="SQLite database path")
    parser.add_argument("--user-id", type=int, default=1, help="Target Logbook user ID (default: 1)")
    parser.add_argument("--append", action="store_true", help="Append instead of resetting the target user's flights/rates")
    args = parser.parse_args()
    import_workbook(args.xlsx, args.db, user_id=args.user_id, reset=not args.append)
