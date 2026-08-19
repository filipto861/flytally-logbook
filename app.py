from __future__ import annotations

import base64
import json
import math
import re
import sqlite3
import xml.etree.ElementTree as ET
from datetime import date, datetime, time, timezone, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urlencode
from zoneinfo import ZoneInfo

import requests

import folium
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
import streamlit.components.v1 as components
try:
    from st_aggrid import AgGrid, DataReturnMode, GridOptionsBuilder, GridUpdateMode
    HAS_AGGRID = True
except Exception:
    AgGrid = None
    DataReturnMode = None
    GridOptionsBuilder = None
    GridUpdateMode = None
    HAS_AGGRID = False
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from streamlit_folium import st_folium

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DB_PATH = DATA_DIR / "logbook.sqlite"
AIRPORT_OVERRIDES_PATH = DATA_DIR / "airport_overrides.csv"
AIRPORTS_CSV_PATH = DATA_DIR / "airports.csv"
AIRPORTS_DB_PATH = DATA_DIR / "airports_full.sqlite"
OURAIRPORTS_AIRPORTS_URL = "https://davidmegginson.github.io/ourairports-data/airports.csv"
APP_VERSION = "v0.35.6"
LOCAL_TZ = ZoneInfo("Europe/Prague")
DB_SCHEMA_VERSION = 3
_DB_READY = False

EVIDENCE_OPTIONS = ["ULL", "EASA"]
CLASS_OPTIONS = ["ULL", "SEP", "TMG", "MEP", "SET", "OTHER", "GLIDER"]
ROLE_OPTIONS = ["PIC", "DUAL", "INSTRUKTOR", "SAFETY PILOT", "CO-PILOT", "PAX", "OBSERVER"]

NAV_ITEMS = [
    ("Dashboard", "Souhrn"),
    ("Lety", "Lety"),
    ("Nový let", "Přidat let"),
    ("Mapa", "Mapa"),
    ("Ceník", "Ceník"),
    ("Databáze", "Databáze"),
    ("Export", "Export"),
]

SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS app_meta (
    key TEXT PRIMARY KEY,
    value TEXT,
    updated_at TEXT
);
CREATE TABLE IF NOT EXISTS flights (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    evidence TEXT,
    registration TEXT,
    aircraft_type TEXT,
    aircraft_class TEXT,
    departure TEXT,
    arrival TEXT,
    off_block TEXT,
    takeoff TEXT,
    landing TEXT,
    on_block TEXT,
    starts INTEGER DEFAULT 1,
    commander TEXT,
    instructor TEXT,
    role TEXT,
    task TEXT,
    price_per_hour REAL,
    note TEXT
);
CREATE TABLE IF NOT EXISTS aircraft (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    registration TEXT NOT NULL UNIQUE,
    aircraft_type TEXT,
    icao_type TEXT,
    aircraft_class TEXT,
    evidence TEXT,
    default_price_per_hour REAL,
    active INTEGER DEFAULT 1,
    note TEXT,
    created_at TEXT,
    updated_at TEXT
);
CREATE TABLE IF NOT EXISTS rates (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    registration TEXT NOT NULL,
    aircraft_type TEXT,
    valid_from TEXT,
    price_per_hour REAL,
    dry_price_per_hour REAL,
    source TEXT,
    UNIQUE(registration, valid_from)
);
CREATE TABLE IF NOT EXISTS airports (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ident TEXT NOT NULL UNIQUE,
    name TEXT,
    airport_type TEXT,
    iso_country TEXT,
    iso_region TEXT,
    municipality TEXT,
    latitude_deg REAL,
    longitude_deg REAL,
    elevation_ft REAL,
    gps_code TEXT,
    iata_code TEXT,
    local_code TEXT,
    source TEXT,
    active INTEGER DEFAULT 1,
    closed INTEGER DEFAULT 0,
    data_quality TEXT,
    imported_at TEXT,
    updated_at TEXT,
    raw_json TEXT
);
CREATE TABLE IF NOT EXISTS flight_tracks (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    flight_id INTEGER NOT NULL,
    file_name TEXT,
    imported_at TEXT,
    point_count INTEGER,
    distance_km REAL,
    start_utc TEXT,
    end_utc TEXT,
    min_alt_m REAL,
    max_alt_m REAL,
    coordinates_json TEXT NOT NULL,
    FOREIGN KEY(flight_id) REFERENCES flights(id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS track_points (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    track_id INTEGER NOT NULL,
    seq INTEGER NOT NULL,
    time_utc TEXT,
    latitude_deg REAL NOT NULL,
    longitude_deg REAL NOT NULL,
    altitude_m REAL,
    segment_km REAL,
    distance_km REAL,
    speed_kmh REAL,
    speed_kt REAL,
    source TEXT,
    FOREIGN KEY(track_id) REFERENCES flight_tracks(id) ON DELETE CASCADE,
    UNIQUE(track_id, seq)
);
CREATE TABLE IF NOT EXISTS audit_log (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    created_at TEXT NOT NULL,
    actor TEXT,
    action TEXT NOT NULL,
    object_type TEXT,
    object_id TEXT,
    detail_json TEXT
);
CREATE INDEX IF NOT EXISTS idx_flights_date ON flights(date);
CREATE INDEX IF NOT EXISTS idx_flights_registration ON flights(registration);
CREATE INDEX IF NOT EXISTS idx_flights_evidence_role ON flights(evidence, role);
CREATE INDEX IF NOT EXISTS idx_tracks_flight_id ON flight_tracks(flight_id);
CREATE INDEX IF NOT EXISTS idx_track_points_track_seq ON track_points(track_id, seq);
CREATE INDEX IF NOT EXISTS idx_airports_ident ON airports(ident);
CREATE INDEX IF NOT EXISTS idx_airports_country ON airports(iso_country);
CREATE INDEX IF NOT EXISTS idx_airports_active ON airports(active, closed);
CREATE INDEX IF NOT EXISTS idx_audit_created_at ON audit_log(created_at);
"""

# -----------------------------------------------------------------------------
# Database
# -----------------------------------------------------------------------------
# -----------------------------------------------------------------------------

def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _set_meta(con: sqlite3.Connection, key: str, value: Any) -> None:
    con.execute(
        "INSERT OR REPLACE INTO app_meta (key, value, updated_at) VALUES (?, ?, ?)",
        (key, str(value), _now_iso()),
    )


def _get_secret(section: str, key: str, default: Any = None) -> Any:
    try:
        sec = st.secrets.get(section, {})
        if hasattr(sec, "get"):
            return sec.get(key, default)
    except Exception:
        pass
    return default


def auth_configured() -> bool:
    return bool(_get_secret("auth", "admin_password", ""))


def is_admin() -> bool:
    return st.session_state.get("auth_role") == "admin"


def actor_name() -> str:
    return "admin" if is_admin() else "viewer"


def invalidate_cached_data() -> None:
    """Clear cached database reads after any write operation."""
    try:
        st.cache_data.clear()
    except Exception:
        pass


def render_auth_sidebar() -> None:
    admin_password = _get_secret("auth", "admin_password", "")
    st.divider()
    with st.expander("Správa aplikace", expanded=is_admin()):
        if admin_password:
            if is_admin():
                if st.button("Odhlásit", use_container_width=True):
                    st.session_state.pop("auth_role", None)
                    st.rerun()
            else:
                with st.form("admin_login_form"):
                    pwd = st.text_input("Heslo správce", type="password")
                    submitted = st.form_submit_button("Přihlásit", use_container_width=True)
                if submitted:
                    if pwd == admin_password:
                        st.session_state["auth_role"] = "admin"
                        st.rerun()
                    else:
                        st.error("Nesprávné heslo")
        else:
            pass


def require_admin() -> bool:
    if is_admin():
        return True
    st.warning("Tato akce je dostupná jen po přihlášení jako admin.")
    return False


def record_audit(con: sqlite3.Connection, action: str, object_type: str | None = None, object_id: Any = None, detail: Any = None) -> None:
    payload = json.dumps(detail, ensure_ascii=False, default=str) if detail is not None else None
    params = (_now_iso(), actor_name(), action, object_type, str(object_id) if object_id is not None else None, payload)
    try:
        con.execute(
            "INSERT INTO audit_log (created_at, actor, action, object_type, object_id, detail_json) VALUES (?, ?, ?, ?, ?, ?)",
            params,
        )
    except sqlite3.OperationalError:
        try:
            ensure_schema_compatibility(con)
            con.execute(
                "INSERT INTO audit_log (created_at, actor, action, object_type, object_id, detail_json) VALUES (?, ?, ?, ?, ?, ?)",
                params,
            )
        except Exception:
            # Audit logging must never block the real database action.
            pass
    except Exception:
        pass
    try:
        _set_meta(con, "last_change_at", _now_iso())
        _set_meta(con, "dirty", "1")
        invalidate_cached_data()
    except Exception:
        pass


def github_backup_configured() -> bool:
    return bool(_get_secret("github", "token", "") or _get_secret("github_sync", "token", ""))


def github_backup_config() -> dict[str, str]:
    token = _get_secret("github", "token", "") or _get_secret("github_sync", "token", "")
    repo = _get_secret("github", "repo", "") or _get_secret("github_sync", "repo", "filipto861/Logbook")
    db_path = _get_secret("github", "db_path", "") or _get_secret("github_sync", "db_path", "data/logbook.sqlite")
    branch = _get_secret("github", "branch", "") or _get_secret("github_sync", "branch", "main")
    auto_backup = _get_secret("github", "auto_backup", "") or _get_secret("github_sync", "auto_backup", "true")
    return {"token": token, "repo": repo, "db_path": db_path, "branch": branch, "auto_backup": str(auto_backup)}


def github_auto_backup_enabled() -> bool:
    cfg = github_backup_config()
    return github_backup_configured() and str(cfg.get("auto_backup", "true")).strip().lower() not in {"0", "false", "no", "off"}


def _record_audit_clean(con: sqlite3.Connection, action: str, object_type: str | None = None, object_id: Any = None, detail: Any = None) -> None:
    """Audit entry that does not mark the database dirty. Used by backup itself."""
    payload = json.dumps(detail, ensure_ascii=False, default=str) if detail is not None else None
    params = (_now_iso(), actor_name(), action, object_type, str(object_id) if object_id is not None else None, payload)
    try:
        con.execute(
            "INSERT INTO audit_log (created_at, actor, action, object_type, object_id, detail_json) VALUES (?, ?, ?, ?, ?, ?)",
            params,
        )
    except sqlite3.OperationalError:
        try:
            ensure_schema_compatibility(con)
            con.execute(
                "INSERT INTO audit_log (created_at, actor, action, object_type, object_id, detail_json) VALUES (?, ?, ?, ?, ?, ?)",
                params,
            )
        except Exception:
            pass


def checkpoint_database() -> None:
    with sqlite3.connect(DB_PATH) as con:
        try:
            con.execute("PRAGMA wal_checkpoint(FULL)")
        except sqlite3.DatabaseError:
            pass


def backup_database_to_github(commit_message: str | None = None) -> str:
    cfg = github_backup_config()
    if not cfg["token"]:
        raise RuntimeError("Chybí GitHub token ve Streamlit secrets.")
    if not DB_PATH.exists():
        raise RuntimeError("Databázový soubor neexistuje.")

    api_url = f"https://api.github.com/repos/{cfg['repo']}/contents/{cfg['db_path']}"
    headers = {"Authorization": f"Bearer {cfg['token']}", "Accept": "application/vnd.github+json"}
    branch = cfg.get("branch") or "main"
    backup_at = _now_iso()

    # Make the database file itself contain the fact that it is backed up. If the
    # remote upload fails, the dirty flag is restored below.
    with connect() as con:
        _set_meta(con, "last_github_backup_at", backup_at)
        _set_meta(con, "last_github_backup_error", "")
        _set_meta(con, "dirty", "0")
        _record_audit_clean(con, "github_backup", "database", cfg["db_path"], {"repo": cfg["repo"], "branch": branch})
        con.commit()

    try:
        checkpoint_database()
        sha = None
        get_resp = requests.get(api_url, headers=headers, params={"ref": branch}, timeout=30)
        if get_resp.status_code == 200:
            sha = get_resp.json().get("sha")
        elif get_resp.status_code not in (404,):
            raise RuntimeError(f"GitHub GET selhal: {get_resp.status_code} {get_resp.text[:300]}")
        content_b64 = base64.b64encode(DB_PATH.read_bytes()).decode("ascii")
        payload = {
            "message": commit_message or f"Backup logbook database {datetime.now(LOCAL_TZ).strftime('%Y-%m-%d %H:%M')}",
            "content": content_b64,
            "branch": branch,
        }
        if sha:
            payload["sha"] = sha
        put_resp = requests.put(api_url, headers=headers, json=payload, timeout=90)
        if put_resp.status_code not in (200, 201):
            raise RuntimeError(f"GitHub PUT selhal: {put_resp.status_code} {put_resp.text[:500]}")
        return put_resp.json().get("commit", {}).get("html_url", "")
    except Exception as exc:
        with connect() as con:
            _set_meta(con, "dirty", "1")
            _set_meta(con, "last_github_backup_error", str(exc)[:500])
            con.commit()
        raise


def auto_backup_after_change(reason: str) -> None:
    """Automatically persist the current SQLite database to GitHub after a confirmed write.

    Streamlit Community Cloud does not preserve local SQLite changes across every
    restart/redeploy. This makes GitHub the versioned persistent backup for the
    single-user online deployment.
    """
    if not github_auto_backup_enabled():
        st.session_state["last_auto_backup_status"] = "not_configured"
        return
    try:
        with st.spinner("Ukládám databázi na GitHub…"):
            url = backup_database_to_github(
                f"Auto backup after {reason} {datetime.now(LOCAL_TZ).strftime('%Y-%m-%d %H:%M')}"
            )
        st.session_state["last_auto_backup_status"] = "ok"
        st.session_state["last_auto_backup_url"] = url
        st.session_state.pop("last_auto_backup_error", None)
    except Exception as exc:
        st.session_state["last_auto_backup_status"] = "error"
        st.session_state["last_auto_backup_error"] = str(exc)



def restore_database_from_upload(uploaded_file) -> None:
    raw = uploaded_file.read()
    if not raw.startswith(b"SQLite format 3"):
        raise RuntimeError("Nahraný soubor nevypadá jako SQLite databáze.")
    backup_path = DB_PATH.with_suffix(".sqlite.before_restore")
    if DB_PATH.exists():
        checkpoint_database()
        backup_path.write_bytes(DB_PATH.read_bytes())
    DB_PATH.write_bytes(raw)
    global _DB_READY
    _DB_READY = False
    with connect() as con:
        record_audit(con, "restore_database", "database", None, {"file": uploaded_file.name})
        con.commit()


def initialize_database(con: sqlite3.Connection) -> None:
    """Idempotent DB bootstrap/migration.

    The application schema lives in SQLite, not in Excel. Existing v0.4 databases are
    upgraded in place: flights/rates/tracks stay untouched, while airport/aircraft
    registries and normalized track_points are added.
    """
    con.execute("PRAGMA foreign_keys = ON")
    con.execute("PRAGMA busy_timeout = 5000")
    try:
        con.execute("PRAGMA journal_mode = WAL")
    except sqlite3.DatabaseError:
        pass
    con.executescript(SCHEMA)
    ensure_schema_compatibility(con)
    _set_meta(con, "schema_version", DB_SCHEMA_VERSION)
    _seed_airports_from_overrides(con)
    _seed_aircraft_from_existing_data(con)
    _backfill_track_points(con)
    con.commit()


def connect() -> sqlite3.Connection:
    global _DB_READY
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    if not _DB_READY:
        initialize_database(con)
        _DB_READY = True
    else:
        con.execute("PRAGMA foreign_keys = ON")
        con.execute("PRAGMA busy_timeout = 5000")
    return con


def _seed_airports_from_overrides(con: sqlite3.Connection) -> None:
    if not AIRPORT_OVERRIDES_PATH.exists():
        return
    try:
        df = pd.read_csv(AIRPORT_OVERRIDES_PATH)
    except Exception:
        return
    if df.empty:
        return
    import_airports_dataframe(con, df, default_source="manual_override", replace_existing=True)


def _seed_aircraft_from_existing_data(con: sqlite3.Connection) -> None:
    rows = con.execute(
        """
        SELECT registration,
               MAX(aircraft_type) AS aircraft_type,
               MAX(aircraft_class) AS aircraft_class,
               MAX(evidence) AS evidence,
               MAX(price_per_hour) AS price_per_hour
        FROM flights
        WHERE registration IS NOT NULL AND TRIM(registration) <> ''
        GROUP BY registration
        """
    ).fetchall()
    now = _now_iso()
    for row in rows:
        con.execute(
            """
            INSERT INTO aircraft (registration, aircraft_type, icao_type, aircraft_class, evidence, default_price_per_hour, active, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, 1, ?, ?)
            ON CONFLICT(registration) DO UPDATE SET
                aircraft_type=COALESCE(excluded.aircraft_type, aircraft.aircraft_type),
                icao_type=COALESCE(excluded.icao_type, aircraft.icao_type),
                aircraft_class=COALESCE(excluded.aircraft_class, aircraft.aircraft_class),
                evidence=COALESCE(excluded.evidence, aircraft.evidence),
                default_price_per_hour=COALESCE(excluded.default_price_per_hour, aircraft.default_price_per_hour),
                updated_at=excluded.updated_at
            """,
            (
                (row["registration"] or "").upper(),
                row["aircraft_type"],
                row["aircraft_type"],
                row["aircraft_class"],
                row["evidence"],
                row["price_per_hour"],
                now,
                now,
            ),
        )


def _backfill_track_points(con: sqlite3.Connection) -> None:
    existing = con.execute("SELECT COUNT(*) FROM track_points").fetchone()[0]
    tracks = con.execute("SELECT id, coordinates_json FROM flight_tracks").fetchall()
    if existing > 0 or not tracks:
        return
    for tr in tracks:
        try:
            points = json.loads(tr["coordinates_json"] or "[]")
        except Exception:
            points = []
        insert_track_points(con, int(tr["id"]), points)


@st.cache_data(show_spinner=False, ttl=30)
def read_table(table: str) -> pd.DataFrame:
    with connect() as con:
        return pd.read_sql_query(f"SELECT * FROM {table}", con)



def _clean_ident(value: Any) -> str:
    text = normalize_text(value)
    return (text or "").upper()


def _to_float(value: Any) -> float | None:
    if value is None or pd.isna(value):
        return None
    try:
        text = str(value).strip().replace(",", ".")
        return float(text) if text else None
    except ValueError:
        return None


def import_airports_dataframe(
    con: sqlite3.Connection,
    df: pd.DataFrame,
    default_source: str,
    replace_existing: bool = True,
) -> int:
    """Import airport-like rows into the airports registry.

    Supported input columns include both OurAirports names and our local override
    names: ident/name/type/airport_type/latitude_deg/lat/longitude_deg/lon.
    """
    if df.empty:
        return 0
    normalized_cols = {str(c).strip().lower(): c for c in df.columns}

    def val(row, *names, default=None):
        for name in names:
            col = normalized_cols.get(name.lower())
            if col is not None:
                return row.get(col)
        return default

    now = _now_iso()
    count = 0
    for _, row in df.iterrows():
        ident = _clean_ident(val(row, "ident", "code", "local_code", "gps_code"))
        name = normalize_text(val(row, "name", "airport_name"))
        lat = _to_float(val(row, "latitude_deg", "lat", "latitude"))
        lon = _to_float(val(row, "longitude_deg", "lon", "longitude"))
        if not ident or lat is None or lon is None:
            continue
        airport_type = normalize_text(val(row, "airport_type", "type", default="small_airport"))
        closed = 1 if str(airport_type or "").lower() in {"closed", "closed_airport"} else int(_to_float(val(row, "closed", default=0)) or 0)
        active = 0 if closed else int(_to_float(val(row, "active", default=1)) or 1)
        source = normalize_text(val(row, "source")) or default_source
        raw = {str(k): (None if pd.isna(v) else v) for k, v in row.to_dict().items()}
        params = (
            ident,
            name,
            airport_type,
            normalize_text(val(row, "iso_country")),
            normalize_text(val(row, "iso_region")),
            normalize_text(val(row, "municipality")),
            lat,
            lon,
            _to_float(val(row, "elevation_ft")),
            normalize_text(val(row, "gps_code")),
            normalize_text(val(row, "iata_code")),
            normalize_text(val(row, "local_code")) or ident,
            source,
            active,
            closed,
            normalize_text(val(row, "data_quality")) or "imported",
            now,
            now,
            json.dumps(raw, ensure_ascii=False),
        )
        if replace_existing:
            con.execute(
                """
                INSERT INTO airports (ident, name, airport_type, iso_country, iso_region, municipality,
                    latitude_deg, longitude_deg, elevation_ft, gps_code, iata_code, local_code,
                    source, active, closed, data_quality, imported_at, updated_at, raw_json)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(ident) DO UPDATE SET
                    name=excluded.name,
                    airport_type=excluded.airport_type,
                    iso_country=excluded.iso_country,
                    iso_region=excluded.iso_region,
                    municipality=excluded.municipality,
                    latitude_deg=excluded.latitude_deg,
                    longitude_deg=excluded.longitude_deg,
                    elevation_ft=excluded.elevation_ft,
                    gps_code=excluded.gps_code,
                    iata_code=excluded.iata_code,
                    local_code=excluded.local_code,
                    source=excluded.source,
                    active=excluded.active,
                    closed=excluded.closed,
                    data_quality=excluded.data_quality,
                    updated_at=excluded.updated_at,
                    raw_json=excluded.raw_json
                """,
                params,
            )
        else:
            con.execute(
                """
                INSERT OR IGNORE INTO airports (ident, name, airport_type, iso_country, iso_region, municipality,
                    latitude_deg, longitude_deg, elevation_ft, gps_code, iata_code, local_code,
                    source, active, closed, data_quality, imported_at, updated_at, raw_json)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                params,
            )
        count += 1
    _set_meta(con, f"airport_import_{default_source}", f"{count} rows")
    return count


def _airport_query(active_only: bool) -> str:
    query = "SELECT * FROM airports"
    if active_only:
        query += " WHERE active = 1 AND closed = 0 AND latitude_deg IS NOT NULL AND longitude_deg IS NOT NULL"
    query += " ORDER BY ident"
    return query


@st.cache_data(show_spinner=False, ttl=300)
def read_airports(active_only: bool = True) -> pd.DataFrame:
    """Return airport registry from the fixed world DB plus local overrides.

    data/airports_full.sqlite is the stable world airport database. The main
    logbook.sqlite keeps only manual overrides/additions, so flights and KML
    tracks are not overwritten when the airport registry is refreshed. When the
    same ident exists in both databases, the local logbook row wins.
    """
    frames: list[pd.DataFrame] = []
    query = _airport_query(active_only)

    if AIRPORTS_DB_PATH.exists():
        try:
            with sqlite3.connect(AIRPORTS_DB_PATH) as airport_con:
                frames.append(pd.read_sql_query(query, airport_con))
        except Exception:
            pass

    try:
        with connect() as con:
            frames.append(pd.read_sql_query(query, con))
    except Exception:
        pass

    if not frames:
        return pd.DataFrame()
    out = pd.concat(frames, ignore_index=True, sort=False)
    if "ident" in out.columns:
        out["ident"] = out["ident"].fillna("").astype(str).str.upper().str.strip()
        out = out[out["ident"] != ""]
        out = out.drop_duplicates(subset=["ident"], keep="last")
        out = out.sort_values("ident")
    return out.reset_index(drop=True)


def import_ourairports_to_database() -> int:
    """Import world airport database into SQLite.

    Prefer bundled data/airports.csv when present. If it is not present,
    fall back to the public OurAirports CSV URL. The application never
    keeps the world airport list hard-coded in Python.
    """
    if AIRPORTS_CSV_PATH.exists():
        df = pd.read_csv(AIRPORTS_CSV_PATH)
        source_label = "OurAirports bundled CSV"
        source_ref = str(AIRPORTS_CSV_PATH.name)
    else:
        df = pd.read_csv(OURAIRPORTS_AIRPORTS_URL)
        source_label = "OurAirports URL"
        source_ref = OURAIRPORTS_AIRPORTS_URL
    with connect() as con:
        count = import_airports_dataframe(con, df, default_source=source_label, replace_existing=True)
        _seed_airports_from_overrides(con)
        _set_meta(con, "ourairports_source", source_ref)
        _set_meta(con, "ourairports_rows", count)
        _set_meta(con, "ourairports_imported_at", _now_iso())
        record_audit(con, "import_ourairports", "airports", None, {"rows": count, "source": source_ref})
        con.commit()
    auto_backup_after_change("import_ourairports")
    return count


def import_airport_csv_upload(uploaded_file) -> int:
    df = pd.read_csv(uploaded_file)
    with connect() as con:
        count = import_airports_dataframe(con, df, default_source="user_csv", replace_existing=True)
        record_audit(con, "import_airport_csv", "airports", None, {"rows": count, "file": getattr(uploaded_file, "name", None)})
        con.commit()
    auto_backup_after_change("import_airport_csv")
    return count


def insert_track_points(con: sqlite3.Connection, track_id: int, points: list[dict[str, Any]]) -> None:
    if not points:
        return
    profile = profile_from_points(points)
    if profile.empty:
        return
    rows = []
    for _, p in profile.iterrows():
        speed_kmh = None if pd.isna(p.get("speed_kmh")) else float(p.get("speed_kmh"))
        rows.append((
            track_id,
            int(p["idx"]),
            p["time_utc"].isoformat() if pd.notna(p.get("time_utc")) else None,
            float(p["lat"]),
            float(p["lon"]),
            None if pd.isna(p.get("alt_m")) else float(p.get("alt_m")),
            float(p.get("seg_km") or 0),
            float(p.get("distance_km") or 0),
            speed_kmh,
            speed_kmh * 0.539957 if speed_kmh is not None else None,
            "kml",
        ))
    con.executemany(
        """
        INSERT OR REPLACE INTO track_points
        (track_id, seq, time_utc, latitude_deg, longitude_deg, altitude_m,
         segment_km, distance_km, speed_kmh, speed_kt, source)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows,
    )

@st.cache_data(show_spinner=False, ttl=30)
def read_track_counts() -> pd.DataFrame:
    with connect() as con:
        return pd.read_sql_query(
            """
            SELECT flight_id, COUNT(*) AS track_count, COALESCE(SUM(distance_km), 0) AS gps_km
            FROM flight_tracks GROUP BY flight_id
            """,
            con,
        )


@st.cache_data(show_spinner=False, ttl=30)
def read_flights() -> pd.DataFrame:
    flights = read_table("flights")
    flights = compute_metrics(flights)
    tracks = read_track_counts()
    if tracks.empty:
        flights["track_count"] = 0
        flights["gps_km"] = 0.0
    else:
        flights = flights.merge(tracks, how="left", left_on="id", right_on="flight_id")
        flights["track_count"] = flights["track_count"].fillna(0).astype(int)
        flights["gps_km"] = flights["gps_km"].fillna(0.0)
    return flights


@st.cache_data(show_spinner=False, ttl=60)
def read_tracks_joined() -> pd.DataFrame:
    with connect() as con:
        return pd.read_sql_query(
            """
            SELECT t.*, f.date, f.evidence, f.registration, f.aircraft_type, f.aircraft_class,
                   f.departure, f.arrival, f.off_block, f.takeoff, f.landing, f.on_block,
                   f.role, f.starts, f.task, f.commander
            FROM flight_tracks t
            JOIN flights f ON f.id = t.flight_id
            ORDER BY f.date, f.off_block, t.id
            """,
            con,
        )


@st.cache_data(show_spinner=False, ttl=60)
def read_tracks_for_flight(flight_id: int) -> pd.DataFrame:
    with connect() as con:
        return pd.read_sql_query("SELECT * FROM flight_tracks WHERE flight_id = ? ORDER BY id", con, params=(flight_id,))


def normalize_text(value: Any) -> str | None:
    if value is None or pd.isna(value):
        return None
    text = str(value).strip()
    return text if text else None


def _columns_for_table(con: sqlite3.Connection, table: str) -> set[str]:
    try:
        return {row[1] for row in con.execute(f"PRAGMA table_info({table})").fetchall()}
    except sqlite3.DatabaseError:
        return set()


def _add_column_if_missing(con: sqlite3.Connection, table: str, column: str, ddl: str) -> None:
    if column not in _columns_for_table(con, table):
        con.execute(f"ALTER TABLE {table} ADD COLUMN {ddl}")


def ensure_schema_compatibility(con: sqlite3.Connection) -> None:
    """Upgrade older SQLite files without losing data.

    CREATE TABLE IF NOT EXISTS does not add new columns to an already existing
    table. Some earlier online databases had an older audit_log shape
    (user/entity/detail). Deleting a flight must not fail just because the audit
    table is old, so we add all columns used by the current application.
    """
    try:
        _add_column_if_missing(con, "audit_log", "actor", "actor TEXT")
        _add_column_if_missing(con, "audit_log", "object_type", "object_type TEXT")
        _add_column_if_missing(con, "audit_log", "object_id", "object_id TEXT")
        _add_column_if_missing(con, "audit_log", "detail_json", "detail_json TEXT")
        _add_column_if_missing(con, "audit_log", "user", "user TEXT")
        _add_column_if_missing(con, "audit_log", "entity", "entity TEXT")
        _add_column_if_missing(con, "audit_log", "entity_id", "entity_id TEXT")
        _add_column_if_missing(con, "audit_log", "detail", "detail TEXT")
    except sqlite3.DatabaseError:
        pass
    try:
        _add_column_if_missing(con, "flight_tracks", "min_alt_m", "min_alt_m REAL")
        _add_column_if_missing(con, "flight_tracks", "max_alt_m", "max_alt_m REAL")
        _add_column_if_missing(con, "flights", "note", "note TEXT")
    except sqlite3.DatabaseError:
        pass


def normalize_date(value: Any) -> str | None:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date().isoformat()
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text or None


def normalize_time(value: Any) -> str | None:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).strftime("%H:%M")
        except ValueError:
            pass
    return text


def parse_time_to_minutes(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "nat"}:
        return None
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            t = datetime.strptime(text, fmt).time()
            return t.hour * 60 + t.minute
        except ValueError:
            pass
    return None


def minutes_diff(start_value: Any, end_value: Any) -> int | None:
    start = parse_time_to_minutes(start_value)
    end = parse_time_to_minutes(end_value)
    if start is None or end is None:
        return None
    return (end - start) % (24 * 60)


def fmt_minutes(minutes: int | float | None) -> str:
    if minutes is None or pd.isna(minutes):
        return ""
    minutes = int(round(float(minutes)))
    return f"{minutes // 60}:{minutes % 60:02d}"


def fmt_money(value: float | int | None) -> str:
    if value is None or pd.isna(value):
        return ""
    return f"{float(value):,.0f} Kč".replace(",", " ")


def compute_metrics(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df.copy()
    out = df.copy()
    out["date_dt"] = pd.to_datetime(out["date"], errors="coerce")
    out["year"] = out["date_dt"].dt.year
    out["block_minutes"] = [minutes_diff(a, b) for a, b in zip(out["off_block"], out["on_block"])]
    out["air_minutes"] = [minutes_diff(a, b) for a, b in zip(out["takeoff"], out["landing"])]
    out["block_hours"] = out["block_minutes"].fillna(0) / 60.0
    out["air_hours"] = out["air_minutes"].fillna(0) / 60.0
    out["price_per_hour"] = pd.to_numeric(out["price_per_hour"], errors="coerce")
    out["cost"] = out["block_hours"] * out["price_per_hour"].fillna(0)
    out["role"] = out["role"].fillna("").str.upper().str.strip()
    out["evidence"] = out["evidence"].fillna("").str.upper().str.strip()
    out["registration"] = out["registration"].fillna("").str.upper().str.strip()
    out["aircraft_class"] = out["aircraft_class"].fillna("").str.upper().str.strip()
    out["starts"] = pd.to_numeric(out["starts"], errors="coerce").fillna(0).astype(int)
    out["block_time"] = out["block_minutes"].apply(fmt_minutes)
    out["air_time"] = out["air_minutes"].apply(fmt_minutes)
    out["cost_label"] = out["cost"].apply(fmt_money)
    return out

# -----------------------------------------------------------------------------
# UI
# -----------------------------------------------------------------------------

def apply_ui_theme(dark_mode: bool) -> None:
    if dark_mode:
        bg = "#06101d"; panel = "#0b182a"; panel2 = "#10233a"; text = "#e6f0fb"; muted = "#92a8c0"; border = "rgba(125,211,252,.18)"; accent = "#38bdf8"; good = "#22c55e"; warn = "#f59e0b"; shadow = "rgba(0,0,0,.40)"
    else:
        bg = "#f5f8fc"; panel = "#ffffff"; panel2 = "#eaf3ff"; text = "#0f172a"; muted = "#475569"; border = "rgba(15,23,42,.12)"; accent = "#0284c7"; good = "#16a34a"; warn = "#d97706"; shadow = "rgba(15,23,42,.10)"
    sidebar_css = """
    :root {--lb-sidebar-width:16.4rem;}
    section[data-testid="stSidebar"], [data-testid="stSidebar"] {
        display:block !important; visibility:visible !important; opacity:1 !important;
        width:var(--lb-sidebar-width) !important; min-width:var(--lb-sidebar-width) !important; max-width:var(--lb-sidebar-width) !important;
        transform:translateX(0) !important; margin-left:0 !important;
        transition:margin-left 320ms cubic-bezier(.22,.61,.36,1), opacity 220ms ease !important;
        will-change:margin-left; z-index:1000 !important;
    }
    [data-testid="stSidebarContent"] {display:block !important; visibility:visible !important; opacity:1 !important;}
    body.lb-sidebar-hidden section[data-testid="stSidebar"],
    body.lb-sidebar-hidden [data-testid="stSidebar"] {
        margin-left:calc(-1 * var(--lb-sidebar-width)) !important;
    }
    [data-testid="stAppViewContainer"] > .main {transition:margin-left 320ms cubic-bezier(.22,.61,.36,1) !important;}
    [data-testid="stAppViewContainer"] .main .block-container {max-width:1500px !important;}
    #lb-sidebar-toggle {
        position:fixed; top:5.20rem; left:calc(var(--lb-sidebar-width) - 2.36rem);
        z-index:2147483647; width:1.92rem; height:1.92rem; border-radius:999px;
        border:1px solid rgba(56,189,248,.38); background:linear-gradient(135deg,rgba(20,43,72,.98),rgba(9,22,39,.98)); color:#dbeafe;
        font-weight:900; font-size:1.02rem; line-height:1; cursor:pointer;
        display:flex; align-items:center; justify-content:center;
        box-shadow:0 10px 28px rgba(0,0,0,.30), inset 0 1px 0 rgba(255,255,255,.06);
        transition:left 320ms cubic-bezier(.22,.61,.36,1), transform 130ms ease, background 140ms ease, border-color 140ms ease, box-shadow 140ms ease;
    }
    #lb-sidebar-toggle:hover {background:linear-gradient(135deg,rgba(31,64,105,.99),rgba(12,29,50,.99)); border-color:rgba(56,189,248,.68); transform:translateY(-1px); box-shadow:0 12px 32px rgba(0,0,0,.34),0 0 0 3px rgba(56,189,248,.06);}
    body.lb-sidebar-hidden #lb-sidebar-toggle {left:.50rem;}
    @media (max-width: 760px) {
        #lb-sidebar-toggle {top:4.35rem; left:calc(var(--lb-sidebar-width) - 2.35rem);}
        body.lb-sidebar-hidden #lb-sidebar-toggle {left:.42rem;}
    }
    """
    st.markdown(f"""
    <style>
    :root {{--bg:{bg};--panel:{panel};--panel2:{panel2};--text:{text};--muted:{muted};--border:{border};--accent:{accent};--good:{good};--warn:{warn};--shadow:{shadow};}}
    /* Streamlit chrome: hide the top header completely. Do NOT rely on the
       native collapsed sidebar button; on Streamlit Cloud it is not rendered
       consistently when the header is hidden. Instead we force the sidebar to
       stay visible and accessible. */
    header[data-testid="stHeader"] {{display:block !important; visibility:visible !important; height:0 !important; min-height:0 !important; background:transparent !important; pointer-events:none !important; z-index:999998 !important; overflow:visible !important;}}
    header[data-testid="stHeader"] * {{pointer-events:none !important;}}
    div[data-testid="stToolbar"], div[data-testid="stDecoration"], div[data-testid="stStatusWidget"], #MainMenu, footer {{display:none !important; visibility:hidden !important; height:0 !important;}}
    .stDeployButton {{display:none !important;}}
    [data-testid="stSidebarHeader"], [data-testid="stSidebarCollapseButton"], [data-testid="collapsedControl"], [data-testid="stSidebarCollapsedControl"],
    section[data-testid="stSidebar"] [data-testid="stSidebarHeader"],
    section[data-testid="stSidebar"] [data-testid="stSidebarHeader"] button,
    section[data-testid="stSidebar"] button[kind="headerNoPadding"],
    section[data-testid="stSidebar"] button[data-testid="baseButton-headerNoPadding"],
    button[title*="sidebar" i], button[aria-label*="sidebar" i],
    button[title*="Collapse" i], button[aria-label*="Collapse" i],
    button[title*="Close" i], button[aria-label*="Close" i] {{display:none !important; visibility:hidden !important; pointer-events:none !important; width:0 !important; height:0 !important; padding:0 !important; margin:0 !important;}}
    {sidebar_css}
    [data-testid="stAppViewContainer"] > .main {{padding-top:0 !important;}}
    [data-testid="stAppViewContainer"] .main .block-container {{padding-top:0 !important; margin-top:0 !important;}}
    .stApp {{background: radial-gradient(circle at 16% 10%, rgba(56,189,248,.16), transparent 24%), radial-gradient(circle at 88% 3%, rgba(34,197,94,.08), transparent 26%), var(--bg); color:var(--text);}}
    [data-testid="stSidebar"] {{background: linear-gradient(180deg, rgba(11,24,42,.98), rgba(6,16,29,.98)); border-right:1px solid var(--border);}}
    [data-testid="stSidebar"] * {{color:#e6f0fb;}}
    .block-container {{padding-top:0 !important; padding-bottom:3rem; max-width:1500px;}}
    h1,h2,h3 {{letter-spacing:-.025em;}}
    .app-title {{display:flex;align-items:center;justify-content:space-between;gap:1rem;padding:1.05rem 1.25rem;border:1px solid var(--border);background:linear-gradient(135deg,rgba(56,189,248,.14),rgba(15,23,42,.04)),var(--panel);border-radius:20px;box-shadow:0 16px 38px var(--shadow);margin-bottom:1rem;}}
    .app-title-main {{font-size:1.55rem;font-weight:850;color:var(--text);line-height:1.1;}}
    .app-title-sub {{font-size:.86rem;color:var(--muted);margin-top:.18rem;}}
    .sidebar-version {{color:var(--muted);font-size:.78rem;margin-top:-.4rem;margin-bottom:1rem;}}
    .app-badge {{font-size:.78rem;font-weight:800;color:#031421;background:linear-gradient(135deg,var(--accent),#a7f3d0);border-radius:999px;padding:.38rem .72rem;white-space:nowrap;}}
    .metric-card {{border:1px solid var(--border);border-radius:18px;padding:1rem 1.05rem;background:linear-gradient(180deg,rgba(255,255,255,.04),transparent),var(--panel);box-shadow:0 12px 30px var(--shadow);min-height:108px;}}
    .metric-label {{color:var(--muted);font-size:.75rem;text-transform:uppercase;letter-spacing:.08em;font-weight:800;}}
    .metric-value {{color:var(--text);font-size:1.72rem;line-height:1.25;font-weight:850;margin-top:.25rem;}}
    .metric-sub {{color:var(--muted);font-size:.82rem;margin-top:.28rem;}}
    #lb-page-loader {{position:fixed;left:var(--lb-sidebar-width);right:0;top:0;bottom:0;z-index:2147483000;display:flex;align-items:flex-start;justify-content:center;padding-top:5.8rem;background:rgba(6,16,29,.10);opacity:0;pointer-events:none;transition:opacity 120ms ease;}}
    body.lb-sidebar-hidden #lb-page-loader {{left:0;}}
    body.lb-page-loading #lb-page-loader {{opacity:1;}}
    .lb-plane-spinner {{width:2.25rem;height:2.25rem;border-radius:999px;border:1px solid rgba(56,189,248,.30);background:rgba(9,22,39,.82);display:flex;align-items:center;justify-content:center;box-shadow:0 10px 28px rgba(0,0,0,.28),0 0 0 4px rgba(56,189,248,.045);color:#a7f3ff;}}
    .lb-plane-spinner span {{display:block;font-size:1.20rem;line-height:1;animation:lbPlaneSpin .78s linear infinite;transform-origin:center center;}}
    @keyframes lbPlaneSpin {{from {{transform:rotate(0deg);}} to {{transform:rotate(360deg);}}}}
    .map-mode-row {{margin:.35rem 0 .65rem 0;}}
    .map-selection-panel {{border:1px solid var(--border);border-radius:18px;background:linear-gradient(180deg,rgba(255,255,255,.035),transparent),var(--panel);padding:.95rem 1.05rem;margin:.8rem 0 1rem 0;box-shadow:0 12px 28px var(--shadow);}}
    .map-selection-title {{font-size:1.06rem;font-weight:900;color:var(--text);letter-spacing:-.015em;margin-bottom:.4rem;}}
    .map-selection-meta {{display:flex;gap:.45rem;flex-wrap:wrap;color:var(--muted);font-size:.82rem;margin-bottom:.65rem;}}
    .map-selection-meta span {{border:1px solid var(--border);border-radius:999px;background:rgba(255,255,255,.035);padding:.16rem .48rem;}}
    .map-mini-head {{color:var(--muted);font-size:.68rem;text-transform:uppercase;letter-spacing:.07em;font-weight:850;padding:.22rem 0;}}
    .map-mini-cell {{border-top:1px solid rgba(148,163,184,.12);padding:.38rem 0;color:var(--text);font-size:.84rem;line-height:1.2;}}
    .map-mini-sub {{color:var(--muted);font-size:.75rem;margin-top:.08rem;}}

    .flight-detail-hero {{border:1px solid var(--border);border-radius:18px;background:linear-gradient(135deg,rgba(56,189,248,.12),rgba(15,23,42,.02)),var(--panel);padding:1rem 1.1rem;margin:.15rem 0 1rem 0;box-shadow:0 12px 28px var(--shadow);}}
    .flight-detail-route {{font-size:1.35rem;font-weight:900;color:var(--text);line-height:1.15;letter-spacing:-.025em;}}
    .flight-detail-meta {{color:var(--muted);font-size:.86rem;margin-top:.35rem;display:flex;gap:.55rem;flex-wrap:wrap;}}
    .flight-detail-meta span {{border:1px solid var(--border);border-radius:999px;background:rgba(255,255,255,.035);padding:.18rem .50rem;}}
    .detail-grid {{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:.75rem;margin:.8rem 0;}}
    .detail-card {{border:1px solid var(--border);border-radius:16px;background:var(--panel);padding:.78rem .85rem;box-shadow:0 10px 24px var(--shadow);min-height:5.4rem;}}
    .detail-card-label {{font-size:.68rem;color:var(--muted);text-transform:uppercase;letter-spacing:.07em;font-weight:850;margin-bottom:.32rem;}}
    .detail-card-value {{font-size:1.08rem;color:var(--text);font-weight:900;line-height:1.1;}}
    .detail-card-sub {{font-size:.78rem;color:var(--muted);margin-top:.25rem;line-height:1.25;}}
    .detail-split {{display:grid;grid-template-columns:1fr 1fr;gap:.75rem;margin:.55rem 0 .9rem 0;}}
    .detail-kv {{border:1px solid var(--border);border-radius:16px;background:rgba(255,255,255,.025);padding:.8rem .9rem;}}
    .detail-kv-title {{font-weight:900;color:var(--text);margin-bottom:.45rem;}}
    .detail-kv-row {{display:flex;justify-content:space-between;gap:1rem;border-top:1px solid rgba(148,163,184,.12);padding:.42rem 0;font-size:.88rem;}}
    .detail-kv-row:first-of-type {{border-top:0;}}
    .detail-kv-row span:first-child {{color:var(--muted);}}
    .detail-kv-row span:last-child {{color:var(--text);font-weight:750;text-align:right;}}
    .section-card {{border:1px solid var(--border);border-radius:18px;padding:1rem;background:var(--panel);box-shadow:0 10px 28px var(--shadow);}}
    .pill {{display:inline-block;border:1px solid var(--border);border-radius:999px;background:var(--panel2);padding:.25rem .62rem;margin:.1rem .18rem;font-size:.82rem;color:var(--text);}}
    div[data-testid="stDataFrame"], div[data-testid="stDataEditor"] {{border-radius:16px;overflow:hidden;}}
    .flight-help {{color:var(--muted);font-size:.86rem;margin:.25rem 0 .75rem 0;}}

    .flight-list-note {{color:var(--muted);font-size:.84rem;margin:.25rem 0 .6rem 0;}}
    .flight-list-head {{font-size:.70rem;color:var(--muted);text-transform:uppercase;letter-spacing:.06em;font-weight:850;padding:.10rem .15rem .20rem .15rem;border-bottom:1px solid var(--border);height:1.35rem;display:flex;align-items:flex-end;}}
    .flight-cell {{font-size:.78rem;line-height:1.05;padding:.06rem .15rem;color:var(--text);min-height:1.78rem;display:flex;flex-direction:column;justify-content:flex-start;}}
    .flight-cell-main {{font-weight:750;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;}}
    .flight-cell-sub {{font-size:.68rem;color:var(--muted);margin-top:.10rem;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;}}
    .flight-row-sep {{height:1px;background:rgba(148,163,184,.10);margin:.04rem 0 .04rem 0;}}
    .flight-page-info {{color:var(--muted);font-size:.82rem;padding-top:1.85rem;text-align:right;}}
    .selected-flight-box {{border:1px solid var(--border); border-radius:14px; padding:.65rem .85rem; background:rgba(56,189,248,.07); margin:.5rem 0 .75rem 0;}}
    .selected-flight-title {{font-weight:850;color:var(--text);}}
    .selected-flight-sub {{font-size:.82rem;color:var(--muted);margin-top:.1rem;}}
    .stTabs [data-baseweb="tab-list"] {{gap:.45rem;}}
    .stTabs [data-baseweb="tab"] {{border-radius:999px;padding:.45rem .9rem;background:var(--panel2);}}
    div.stButton > button {{border-radius:14px !important; font-weight:800 !important; border:1px solid var(--border) !important; min-height:2.05rem; padding:.22rem .60rem !important;}}
    .stButton {{margin-top:0 !important;}}
    [data-testid="column"] .stButton > button {{min-height:1.72rem !important;padding:.08rem .42rem !important;border-radius:11px !important;font-size:.78rem !important;}}
    div.stButton > button[kind="primary"] {{box-shadow:0 10px 22px rgba(56,189,248,.18) !important;}}
    div[data-testid="stExpander"] {{border:1px solid var(--border); border-radius:16px; background:rgba(255,255,255,.025);}}
    div[data-testid="stDialog"] div[role="dialog"] {{border:1px solid var(--border); border-radius:22px;}}
    button[kind="primary"] {{border-radius:12px;}}
    @media (max-width: 980px) {{.detail-grid {{grid-template-columns:repeat(2,minmax(0,1fr));}} .detail-split {{grid-template-columns:1fr;}}}}
    @media (max-width: 760px) {{.block-container {{padding-left:.75rem;padding-right:.75rem;}} .app-title {{padding:.85rem;border-radius:15px;}} .app-title-main {{font-size:1.2rem;}} .metric-value {{font-size:1.35rem;}} .detail-grid {{grid-template-columns:1fr;}}}}
    </style>
    """, unsafe_allow_html=True)


def app_header(subtitle: str = "Osobní letový zápisník") -> None:
    st.markdown(f"""
    <div class="app-title">
      <div><div class="app-title-main">Letový zápisník</div><div class="app-title-sub">{subtitle}</div></div>
      <div class="app-badge">{APP_VERSION}</div>
    </div>
    """, unsafe_allow_html=True)


def metric_card(label: str, value: str, sub: str = "") -> None:
    st.markdown(f"""
    <div class="metric-card"><div class="metric-label">{label}</div><div class="metric-value">{value}</div><div class="metric-sub">{sub}</div></div>
    """, unsafe_allow_html=True)


def go_to_page(page: str) -> None:
    st.session_state["page"] = page
    st.rerun()


def render_nav_button(page_name: str, label: str, key: str) -> None:
    current = st.session_state.get("page", "Dashboard") == page_name
    if st.button(label, key=key, use_container_width=True, type="primary" if current else "secondary"):
        go_to_page(page_name)


def render_sidebar_nav() -> None:
    st.markdown("### Navigace")
    for page_name, label in NAV_ITEMS:
        render_nav_button(page_name, label, f"side_nav_{page_name}")


def get_selected_dataframe_rows(event: Any) -> list[int]:
    if event is None:
        return []
    try:
        return list(event.selection.rows)
    except Exception:
        pass
    try:
        return list(event["selection"]["rows"])
    except Exception:
        return []


def clear_open_flight_dialog() -> None:
    current = st.session_state.get("open_flight_dialog_id")
    if current is not None:
        st.session_state["dismissed_flight_id"] = current
    st.session_state.pop("open_flight_dialog_id", None)
    try:
        if "flight_id" in st.query_params:
            del st.query_params["flight_id"]
    except Exception:
        pass


def _query_param_value(name: str) -> str | None:
    try:
        value = st.query_params.get(name)
        if isinstance(value, list):
            return value[0] if value else None
        return value
    except Exception:
        return None


def base_app_url() -> str:
    try:
        current_url = str(getattr(st.context, "url", "") or "")
        if current_url.startswith(("http://", "https://")):
            return current_url.split("?")[0].split("#")[0]
    except Exception:
        pass
    return ""


def app_link(**params: Any) -> str:
    clean = {str(k): str(v) for k, v in params.items() if v is not None and str(v) != ""}
    qs = urlencode(clean)
    base = base_app_url()
    if base:
        return f"{base}?{qs}" if qs else base
    return f"?{qs}" if qs else "?"


def detail_link(flight_id: int) -> str:
    return app_link(flight_id=int(flight_id))


def airport_link(ident: str) -> str:
    return app_link(map_airport=str(ident).upper())


def route_link(dep: str, arr: str) -> str:
    return app_link(map_route=f"{str(dep).upper()}__{str(arr).upper()}")


def plotly_layout(fig):
    fig.update_layout(template="plotly_dark", paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", font=dict(color="#e5edf7"), margin=dict(l=10, r=10, t=45, b=10), legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1))
    fig.update_xaxes(gridcolor="rgba(148,163,184,.18)")
    fig.update_yaxes(gridcolor="rgba(148,163,184,.18)")
    return fig

# -----------------------------------------------------------------------------
# Filters and summaries
# -----------------------------------------------------------------------------

def apply_filters(df: pd.DataFrame, key_prefix: str = "") -> pd.DataFrame:
    if df.empty:
        return df
    work = df.copy()
    years = sorted(int(y) for y in work["year"].dropna().unique())
    registrations = sorted(r for r in work["registration"].dropna().unique() if r)
    roles = sorted(r for r in work["role"].dropna().unique() if r)

    with st.expander("Filtry", expanded=False):
        f1, f2, f3, f4 = st.columns(4)
        with f1:
            selected_years = st.multiselect("Rok", years, default=years, key=f"{key_prefix}_years")
        with f2:
            selected_evidence = st.multiselect("Evidence", EVIDENCE_OPTIONS, default=EVIDENCE_OPTIONS, key=f"{key_prefix}_ev")
        with f3:
            selected_roles = st.multiselect("Funkce", roles, default=roles, key=f"{key_prefix}_roles")
        with f4:
            selected_regs = st.multiselect("Imatrikulace", registrations, default=[], key=f"{key_prefix}_regs")

    if selected_years:
        work = work[work["year"].isin(selected_years)]
    if selected_evidence:
        work = work[work["evidence"].isin(selected_evidence)]
    if selected_roles:
        work = work[work["role"].isin(selected_roles)]
    if selected_regs:
        work = work[work["registration"].isin(selected_regs)]
    return work


def stat_minutes(df: pd.DataFrame, mask=None) -> int:
    if df.empty:
        return 0
    series = df["block_minutes"].fillna(0)
    if mask is not None:
        series = series[mask]
    return int(series.sum())


def build_summary(df: pd.DataFrame) -> dict[str, int | float]:
    if df.empty:
        return {"flights":0,"starts":0,"total":0,"air":0,"pic":0,"pic_ull":0,"pic_easa":0,"dual":0,"safety":0,"ull":0,"easa":0,"cost":0.0,"tracks":0,"gps_km":0.0}
    return {
        "flights": int(len(df)),
        "starts": int(df["starts"].sum()),
        "total": int(df["block_minutes"].fillna(0).sum()),
        "air": int(df["air_minutes"].fillna(0).sum()),
        "pic": stat_minutes(df, df["role"].eq("PIC")),
        "pic_ull": stat_minutes(df, df["role"].eq("PIC") & df["evidence"].eq("ULL")),
        "pic_easa": stat_minutes(df, df["role"].eq("PIC") & df["evidence"].eq("EASA")),
        "dual": stat_minutes(df, df["role"].eq("DUAL")),
        "safety": stat_minutes(df, df["role"].eq("SAFETY PILOT")),
        "ull": stat_minutes(df, df["evidence"].eq("ULL")),
        "easa": stat_minutes(df, df["evidence"].eq("EASA")),
        "cost": float(df["cost"].fillna(0).sum()),
        "tracks": int(df.get("track_count", pd.Series(dtype=int)).fillna(0).sum()) if "track_count" in df else 0,
        "gps_km": float(df.get("gps_km", pd.Series(dtype=float)).fillna(0).sum()) if "gps_km" in df else 0.0,
    }

# -----------------------------------------------------------------------------
# KML/GPS
# -----------------------------------------------------------------------------

def local_name(tag: str) -> str:
    return tag.split("}", 1)[-1] if "}" in tag else tag


def parse_iso(value: str | None) -> datetime | None:
    if not value:
        return None
    text = value.strip().replace("Z", "+00:00")
    try:
        dt = datetime.fromisoformat(text)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt
    except ValueError:
        return None



def normalize_track_points(points: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Return points in a safe chronological order when timestamps are available.

    Some KML exports store points newest-first. That breaks computed GPS speed
    because time deltas become negative. We keep non-timed tracks in original
    order, but for timed tracks we sort by timestamp and remove clearly invalid
    duplicate coordinates/timestamps only where needed for calculations.
    """
    if not points:
        return []
    indexed = []
    timed_count = 0
    for i, pt in enumerate(points):
        dt = parse_iso(pt.get("time"))
        if dt is not None:
            timed_count += 1
        indexed.append((i, dt, pt))
    if timed_count >= 2:
        indexed.sort(key=lambda item: (item[1] is None, item[1] or datetime.max.replace(tzinfo=timezone.utc), item[0]))
        return [dict(pt) for _, _, pt in indexed]
    return [dict(pt) for pt in points]

def _coord_tokens_to_points(coord_text: str) -> list[dict[str, Any]]:
    parsed: list[dict[str, Any]] = []
    for token in coord_text.replace("\n", " ").replace("\t", " ").split():
        parts = token.split(",")
        if len(parts) < 2:
            continue
        try:
            lon = float(parts[0])
            lat = float(parts[1])
            alt = float(parts[2]) if len(parts) >= 3 and parts[2] else None
        except ValueError:
            continue
        if -90 <= lat <= 90 and -180 <= lon <= 180:
            parsed.append({"lat": lat, "lon": lon, "alt": alt, "time": None})
    return parsed


def _placemark_text(placemark: ET.Element, element_name: str) -> str | None:
    for elem in placemark.iter():
        if local_name(elem.tag) == element_name and (elem.text or "").strip():
            return (elem.text or "").strip()
    return None


def _parse_fr24_description_times(text: str | None) -> list[str]:
    """Return UTC ISO strings found in a Flightradar24 Placemark description/name."""
    if not text:
        return []
    found = re.findall(r"(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2}:\d{2})\s*UTC", text)
    return [f"{date}T{clock}+00:00" for date, clock in found]


def _append_point_unique(points: list[dict[str, Any]], point: dict[str, Any]) -> None:
    if points:
        prev = points[-1]
        same_pos = abs(float(prev.get("lat", 999)) - float(point.get("lat", -999))) < 1e-7 and abs(float(prev.get("lon", 999)) - float(point.get("lon", -999))) < 1e-7
        same_time = (prev.get("time") or None) == (point.get("time") or None)
        if same_pos and same_time:
            return
    points.append(point)


def parse_kml_bytes(data: bytes) -> list[dict[str, Any]]:
    root = ET.fromstring(data)
    points: list[dict[str, Any]] = []

    # Preferred format: gx:Track / Track with <when> + <gx:coord>.
    for track in root.iter():
        if local_name(track.tag) != "Track":
            continue
        whens: list[str | None] = []
        coords: list[str] = []
        for child in list(track):
            lname = local_name(child.tag)
            if lname == "when":
                whens.append((child.text or "").strip() or None)
            elif lname == "coord":
                coords.append((child.text or "").strip())
        for idx, coord in enumerate(coords):
            parts = coord.replace(",", " ").split()
            if len(parts) < 2:
                continue
            try:
                lon = float(parts[0]); lat = float(parts[1]); alt = float(parts[2]) if len(parts) >= 3 else None
            except ValueError:
                continue
            if -90 <= lat <= 90 and -180 <= lon <= 180:
                points.append({"lat": lat, "lon": lon, "alt": alt, "time": whens[idx] if idx < len(whens) else None})
    if points:
        return normalize_track_points(points)

    # Flightradar24 usually exports two parallel datasets in one KML:
    # 1) timestamped Point Placemarks,
    # 2) visual LineString segments named P-1, P-2, ...
    # Older parser mixed both together, which produced an artificial line from
    # the last real point back to the first segment. Therefore Point Placemarks
    # are parsed first and, when they contain usable times, the LineStrings are
    # intentionally ignored.
    point_points: list[dict[str, Any]] = []
    line_segment_points: list[dict[str, Any]] = []

    for placemark in root.iter():
        if local_name(placemark.tag) != "Placemark":
            continue

        pm_name = _placemark_text(placemark, "name") or ""
        pm_desc = _placemark_text(placemark, "description") or ""
        when = _placemark_text(placemark, "when") or _placemark_text(placemark, "begin")

        # Point-only placemarks: FR24 exports one of these for every recorded fix.
        for point_elem in placemark.iter():
            if local_name(point_elem.tag) != "Point":
                continue
            coord_text = None
            for elem in point_elem.iter():
                if local_name(elem.tag) == "coordinates" and (elem.text or "").strip():
                    coord_text = (elem.text or "").strip()
                    break
            if not coord_text:
                continue
            parsed = _coord_tokens_to_points(coord_text)
            if not parsed:
                continue
            pt = parsed[0]
            # Some FR24 files put the timestamp in the Placemark name instead of <when>.
            time_from_name = _parse_fr24_description_times(pm_name)
            pt["time"] = when or (time_from_name[0] if time_from_name else None)
            _append_point_unique(point_points, pt)

        # LineString segments: keep them only as a fallback when no timestamped
        # Point stream exists. Parse the two timestamps from the description.
        for line_elem in placemark.iter():
            if local_name(line_elem.tag) != "LineString":
                continue
            coord_text = None
            for elem in line_elem.iter():
                if local_name(elem.tag) == "coordinates" and (elem.text or "").strip():
                    coord_text = (elem.text or "").strip()
                    break
            if not coord_text:
                continue
            parsed = _coord_tokens_to_points(coord_text)
            if not parsed:
                continue
            times = _parse_fr24_description_times(pm_desc)
            for idx, pt in enumerate(parsed):
                if idx < len(times):
                    pt["time"] = times[idx]
                _append_point_unique(line_segment_points, pt)

    timed_points = [p for p in point_points if p.get("time")]
    if len(timed_points) >= 2:
        return normalize_track_points(point_points)
    if len(line_segment_points) >= 2:
        return normalize_track_points(line_segment_points)
    if point_points:
        return normalize_track_points(point_points)

    # Fallback: plain coordinates. This is intentionally only reached when no
    # usable Point/LineString track was parsed above.
    for elem in root.iter():
        if local_name(elem.tag) != "coordinates":
            continue
        text = (elem.text or "").strip()
        for pt in _coord_tokens_to_points(text):
            _append_point_unique(points, pt)
    return normalize_track_points(points)


def haversine_km(a: dict[str, Any], b: dict[str, Any]) -> float:
    r = 6371.0088
    lat1 = math.radians(float(a["lat"])); lat2 = math.radians(float(b["lat"]))
    dlat = lat2 - lat1; dlon = math.radians(float(b["lon"]) - float(a["lon"]))
    h = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return 2 * r * math.asin(math.sqrt(h))


def track_distance_km(points: list[dict[str, Any]]) -> float:
    if len(points) < 2:
        return 0.0
    return sum(haversine_km(a, b) for a, b in zip(points[:-1], points[1:]))


def track_stats(points: list[dict[str, Any]]) -> dict[str, Any]:
    times = [p.get("time") for p in points if p.get("time")]
    alts = [float(p["alt"]) for p in points if p.get("alt") is not None and float(p.get("alt") or 0) != 0]
    return {
        "point_count": len(points),
        "distance_km": track_distance_km(points),
        "start_utc": times[0] if times else None,
        "end_utc": times[-1] if times else None,
        "min_alt_m": min(alts) if alts else None,
        "max_alt_m": max(alts) if alts else None,
    }


def profile_from_points(points: list[dict[str, Any]], tz: ZoneInfo = LOCAL_TZ) -> pd.DataFrame:
    points = normalize_track_points(points)
    rows = []
    cum = 0.0
    prev = None
    prev_dt = None
    for i, p in enumerate(points):
        dt_utc = parse_iso(p.get("time"))
        dt_local = dt_utc.astimezone(tz) if dt_utc else None
        seg_km = haversine_km(prev, p) if prev is not None else 0.0
        cum += seg_km

        # Prefer speed already present in a future/imported JSON, otherwise compute
        # groundspeed from distance and timestamps. This keeps old and new tracks
        # compatible and fixes KML files exported in reverse chronological order.
        speed = None
        for key in ("speed_kmh", "gps_speed_kmh", "speed"):
            raw = p.get(key)
            if raw is not None:
                try:
                    speed = float(raw)
                    break
                except (TypeError, ValueError):
                    pass
        if speed is None and prev is not None and dt_utc is not None and prev_dt is not None:
            seconds = (dt_utc - prev_dt).total_seconds()
            if seconds > 0:
                speed = seg_km / (seconds / 3600)
                # Reject impossible spikes from malformed KML/timestamp jumps.
                if speed > 900:
                    speed = None

        alt_m = None if p.get("alt") is None else float(p.get("alt"))
        rows.append({
            "idx": i,
            "time_utc": dt_utc,
            "time_local": dt_local,
            "lat": float(p["lat"]),
            "lon": float(p["lon"]),
            "alt_m": alt_m,
            "alt_ft": alt_m * 3.28084 if alt_m is not None else None,
            "seg_km": seg_km,
            "distance_km": cum,
            "speed_kmh": speed,
        })
        prev = p
        prev_dt = dt_utc
    df = pd.DataFrame(rows)
    if not df.empty:
        speed_series = pd.to_numeric(df["speed_kmh"], errors="coerce")
        if speed_series.notna().any():
            df["speed_smooth"] = speed_series.interpolate(limit_direction="both").rolling(5, min_periods=1, center=True).median()
        else:
            df["speed_smooth"] = pd.NA
    return df


def _longest_true_segment(mask: pd.Series, min_len: int = 1) -> tuple[int, int] | None:
    best: tuple[int, int] | None = None
    best_len = 0
    start: int | None = None
    values = [bool(v) for v in mask.fillna(False).tolist()]
    for i, val in enumerate(values + [False]):
        if val and start is None:
            start = i
        elif not val and start is not None:
            end = i - 1
            length = end - start + 1
            if length >= min_len and length > best_len:
                best = (start, end)
                best_len = length
            start = None
    return best


def detect_takeoff_landing(points: list[dict[str, Any]]) -> dict[str, Any]:
    prof = profile_from_points(points)
    n = len(prof)
    if n == 0:
        return {}

    speed = prof.get("speed_smooth", pd.Series(dtype=float)).fillna(0)
    moving = speed.gt(12)

    # Primary logic: airborne segment by GPS groundspeed.
    # Use the longest continuous segment so short GPS spikes do not become takeoff/landing.
    airborne = speed.gt(55)
    segment = _longest_true_segment(airborne, min_len=3)

    # Fallback when speed cannot be calculated: use altitude above the lowest recorded level.
    if segment is None and "alt_m" in prof and prof["alt_m"].notna().any():
        alt = prof["alt_m"]
        base_alt = float(alt.dropna().quantile(0.05))
        airborne_alt = alt.gt(base_alt + 35)
        segment = _longest_true_segment(airborne_alt, min_len=3)

    # Final fallback: treat the useful moving part as the flight.
    if segment is None:
        segment = _longest_true_segment(moving, min_len=1)

    first_move = int(moving.idxmax()) if moving.any() else 0
    last_move = int(moving[moving].index[-1]) if moving.any() else n - 1

    if segment is None:
        takeoff, landing = 0, n - 1
    else:
        takeoff, landing = segment

    if landing < takeoff:
        takeoff, landing = 0, n - 1

    return {"off_idx": first_move, "takeoff_idx": int(takeoff), "landing_idx": int(landing), "on_idx": last_move}


def point_local_dt(points: list[dict[str, Any]], idx: int) -> datetime | None:
    if not points:
        return None
    idx = max(0, min(idx, len(points) - 1))
    dt = parse_iso(points[idx].get("time"))
    return dt.astimezone(LOCAL_TZ) if dt else None


def dt_hhmm(dt: datetime | None) -> str | None:
    return dt.strftime("%H:%M") if dt else None


def point_local_hhmm(points: list[dict[str, Any]], idx: int) -> str | None:
    return dt_hhmm(point_local_dt(points, idx))


def inferred_clock_times(points: list[dict[str, Any]], idx: dict[str, Any], block_padding_minutes: int = 5) -> dict[str, str | None]:
    takeoff_dt = point_local_dt(points, idx.get("takeoff_idx", 0))
    landing_dt = point_local_dt(points, idx.get("landing_idx", len(points) - 1))
    off_dt = takeoff_dt - timedelta(minutes=block_padding_minutes) if takeoff_dt else None
    on_dt = landing_dt + timedelta(minutes=block_padding_minutes) if landing_dt else None
    return {
        "off_block": dt_hhmm(off_dt),
        "takeoff": dt_hhmm(takeoff_dt),
        "landing": dt_hhmm(landing_dt),
        "on_block": dt_hhmm(on_dt),
    }


def point_local_date(points: list[dict[str, Any]]) -> date:
    for p in points:
        dt = parse_iso(p.get("time"))
        if dt:
            return dt.astimezone(LOCAL_TZ).date()
    return date.today()


def nearest_airport(point: dict[str, Any] | None, max_km: float = 18.0) -> str:
    if not point:
        return ""
    airports = read_airports(active_only=True)
    if airports.empty:
        return ""
    p_lat = float(point["lat"])
    p_lon = float(point["lon"])
    lat = pd.to_numeric(airports["latitude_deg"], errors="coerce")
    lon = pd.to_numeric(airports["longitude_deg"], errors="coerce")
    valid = lat.notna() & lon.notna()
    if not valid.any():
        return ""
    work = airports.loc[valid].copy()
    lat = lat.loc[valid]
    lon = lon.loc[valid]
    # Vectorized haversine. Fast enough for full OurAirports database.
    r = 6371.0088
    lat1 = math.radians(p_lat)
    lat2 = lat.apply(math.radians)
    dlat = lat2 - lat1
    dlon = (lon - p_lon).apply(math.radians)
    h = (dlat / 2).apply(math.sin).pow(2) + math.cos(lat1) * lat2.apply(math.cos) * (dlon / 2).apply(math.sin).pow(2)
    dist = 2 * r * h.apply(lambda x: math.asin(math.sqrt(min(max(float(x), 0.0), 1.0))))
    best_pos = int(dist.idxmin())
    best_dist = float(dist.loc[best_pos])
    if best_dist <= max_km:
        return str(work.loc[best_pos, "ident"]).upper()
    return ""


def extract_registration_from_filename(name: str) -> str:
    text = name.upper().replace("_", "-").replace(" ", "-")
    m = re.search(r"OK-?[A-Z]{3}\d{2}", text)
    if m:
        val = m.group(0).replace("OK", "OK-").replace("OK--", "OK-")
        return val if val.startswith("OK-") else "OK-" + val[2:]
    m = re.search(r"OK-?[A-Z]{3}(?!\d)", text)
    if m:
        val = m.group(0).replace("OK", "OK-").replace("OK--", "OK-")
        return val if val.startswith("OK-") else "OK-" + val[2:]
    return ""


def evidence_from_registration(reg: str) -> str:
    return "ULL" if re.fullmatch(r"OK-[A-Z]{3}\d{2}", (reg or "").upper()) else "EASA"


def default_class_for(evidence: str) -> str:
    return "ULL" if evidence == "ULL" else "SEP"


def lookup_latest_rate(rates: pd.DataFrame, registration: str) -> dict[str, Any]:
    if rates.empty or not registration:
        return {}
    reg = registration.strip().upper()
    sub = rates[rates["registration"].fillna("").str.upper().eq(reg)].copy()
    if sub.empty:
        return {}
    sub["valid_from_dt"] = pd.to_datetime(sub["valid_from"], errors="coerce")
    row = sub.sort_values("valid_from_dt").iloc[-1]
    return {"aircraft_type": row.get("aircraft_type"), "price_per_hour": row.get("price_per_hour")}


def infer_from_track(points: list[dict[str, Any]], file_name: str, rates: pd.DataFrame) -> dict[str, Any]:
    idx = detect_takeoff_landing(points)
    stats = track_stats(points)
    reg = extract_registration_from_filename(file_name)
    evidence = evidence_from_registration(reg) if reg else "ULL"
    rate = lookup_latest_rate(rates, reg)
    clock = inferred_clock_times(points, idx, block_padding_minutes=5)
    has_clock = any(p.get("time") for p in points)
    note = "" if has_clock else "KML neobsahovalo časové značky, časy je nutné doplnit ručně."
    return {
        "date": point_local_date(points),
        "registration": reg,
        "evidence": evidence,
        "aircraft_type": normalize_text(rate.get("aircraft_type")) or "",
        "aircraft_class": default_class_for(evidence),
        "departure": nearest_airport(points[idx.get("takeoff_idx", idx.get("off_idx", 0))] if points else None),
        "arrival": nearest_airport(points[idx.get("landing_idx", idx.get("on_idx", len(points)-1))] if points else None),
        "off_block": clock.get("off_block"),
        "takeoff": clock.get("takeoff"),
        "landing": clock.get("landing"),
        "on_block": clock.get("on_block"),
        "starts": 1,
        "commander": "Točík Filip",
        "instructor": "",
        "role": "PIC",
        "task": "",
        "price_per_hour": float(rate.get("price_per_hour")) if rate and pd.notna(rate.get("price_per_hour")) else 0.0,
        "note": note,
        "stats": stats,
        "detect_idx": idx,
        "has_clock": has_clock,
    }


def save_track(flight_id: int, file_name: str, points: list[dict[str, Any]], replace_existing: bool = False) -> None:
    points = normalize_track_points(points)
    stats = track_stats(points)
    with connect() as con:
        if replace_existing:
            con.execute("DELETE FROM flight_tracks WHERE flight_id = ?", (flight_id,))
        cur = con.execute(
            """
            INSERT INTO flight_tracks (flight_id, file_name, imported_at, point_count, distance_km, start_utc, end_utc, min_alt_m, max_alt_m, coordinates_json)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (flight_id, file_name, _now_iso(), stats["point_count"], stats["distance_km"], stats["start_utc"], stats["end_utc"], stats["min_alt_m"], stats["max_alt_m"], json.dumps(points, ensure_ascii=False)),
        )
        track_id = int(cur.lastrowid)
        insert_track_points(con, track_id, points)
        record_audit(con, "save_track", "flight_tracks", track_id, {"flight_id": flight_id, "file_name": file_name, "replace_existing": replace_existing})
        con.commit()
    auto_backup_after_change("save_track")


def create_flight(data: dict[str, Any], auto_backup: bool = True) -> int:
    fields = ["date","evidence","registration","aircraft_type","aircraft_class","departure","arrival","off_block","takeoff","landing","on_block","starts","commander","instructor","role","task","price_per_hour","note"]
    values = []
    for f in fields:
        v = data.get(f)
        if f == "date":
            v = normalize_date(v)
        elif f in {"off_block", "takeoff", "landing", "on_block"}:
            v = normalize_time(v)
        elif f == "starts":
            v = int(v or 0)
        elif f == "price_per_hour":
            v = float(v or 0)
        elif f in {"evidence", "registration", "aircraft_class", "departure", "arrival", "role"}:
            v = normalize_text(v)
            v = v.upper() if v else None
        else:
            v = normalize_text(v)
        values.append(v)
    with connect() as con:
        cur = con.execute(
            """
            INSERT INTO flights (date, evidence, registration, aircraft_type, aircraft_class, departure, arrival, off_block, takeoff, landing, on_block, starts, commander, instructor, role, task, price_per_hour, note)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            values,
        )
        flight_id = int(cur.lastrowid)
        record_audit(con, "create_flight", "flights", flight_id, data)
        con.commit()
    if auto_backup:
        auto_backup_after_change("create_flight")
    return flight_id


def update_flight(flight_id: int, data: dict[str, Any]) -> None:
    fields = ["date","evidence","registration","aircraft_type","aircraft_class","departure","arrival","off_block","takeoff","landing","on_block","starts","commander","instructor","role","task","price_per_hour","note"]
    values = []
    for f in fields:
        v = data.get(f)
        if f == "date":
            v = normalize_date(v)
        elif f in {"off_block", "takeoff", "landing", "on_block"}:
            v = normalize_time(v)
        elif f == "starts":
            v = int(v or 0)
        elif f == "price_per_hour":
            v = float(v or 0)
        elif f in {"evidence", "registration", "aircraft_class", "departure", "arrival", "role"}:
            v = normalize_text(v)
            v = v.upper() if v else None
        else:
            v = normalize_text(v)
        values.append(v)
    with connect() as con:
        con.execute("UPDATE flights SET " + ", ".join(f"{f}=?" for f in fields) + " WHERE id=?", [*values, flight_id])
        record_audit(con, "update_flight", "flights", flight_id, data)
        con.commit()
    auto_backup_after_change("update_flight")


def delete_track(track_id: int) -> None:
    with connect() as con:
        con.execute("DELETE FROM flight_tracks WHERE id = ?", (track_id,))
        record_audit(con, "delete_track", "flight_tracks", track_id, None)
        con.commit()
    auto_backup_after_change("delete_track")


def delete_flight(flight_id: int) -> None:
    """Delete one flight and all related KML/GPS data from SQLite."""
    with connect() as con:
        ensure_schema_compatibility(con)
        row = con.execute("SELECT * FROM flights WHERE id = ?", (flight_id,)).fetchone()
        if row is None:
            return
        track_rows = con.execute("SELECT id FROM flight_tracks WHERE flight_id = ?", (flight_id,)).fetchall()
        track_ids = [int(r["id"]) for r in track_rows]
        audit_detail = {
            "date": row["date"],
            "registration": row["registration"],
            "route": f"{row['departure'] or ''}-{row['arrival'] or ''}",
            "tracks_deleted": len(track_ids),
        }
        for track_id in track_ids:
            con.execute("DELETE FROM track_points WHERE track_id = ?", (track_id,))
        con.execute("DELETE FROM flight_tracks WHERE flight_id = ?", (flight_id,))
        con.execute("DELETE FROM flights WHERE id = ?", (flight_id,))
        record_audit(con, "delete_flight", "flights", flight_id, audit_detail)
        con.commit()
    auto_backup_after_change("delete_flight")


def downsample_points(points: list[dict[str, Any]], max_points: int = 1200) -> list[dict[str, Any]]:
    if len(points) <= max_points:
        return points
    step = max(1, math.ceil(len(points) / max_points))
    sampled = points[::step]
    if sampled[-1] != points[-1]:
        sampled.append(points[-1])
    return sampled


@st.cache_data(show_spinner=False, ttl=600)
def airport_coord_lookup() -> dict[str, dict[str, Any]]:
    """Fast airport coordinate lookup used by maps and track extensions.

    This intentionally reads only the five columns needed for drawing maps. The
    full airport registry has tens of thousands of rows and many columns; loading
    it here would make every first map render noticeably slower.
    """
    query = """
        SELECT ident, name, latitude_deg, longitude_deg, source
        FROM airports
        WHERE latitude_deg IS NOT NULL AND longitude_deg IS NOT NULL
    """
    frames: list[pd.DataFrame] = []
    if AIRPORTS_DB_PATH.exists():
        try:
            with sqlite3.connect(AIRPORTS_DB_PATH) as airport_con:
                frames.append(pd.read_sql_query(query, airport_con))
        except Exception:
            pass
    try:
        with connect() as con:
            frames.append(pd.read_sql_query(query, con))
    except Exception:
        pass
    if not frames:
        return {}
    airports = pd.concat(frames, ignore_index=True, sort=False)
    if airports.empty or "ident" not in airports.columns:
        return {}
    airports["ident"] = airports["ident"].fillna("").astype(str).str.upper().str.strip()
    airports = airports[airports["ident"] != ""]
    airports["latitude_deg"] = pd.to_numeric(airports["latitude_deg"], errors="coerce")
    airports["longitude_deg"] = pd.to_numeric(airports["longitude_deg"], errors="coerce")
    airports = airports.dropna(subset=["latitude_deg", "longitude_deg"])
    airports = airports[
        airports["latitude_deg"].between(-90, 90)
        & airports["longitude_deg"].between(-180, 180)
    ]
    airports = airports.drop_duplicates(subset=["ident"], keep="last")
    lookup: dict[str, dict[str, Any]] = {}
    for row in airports.itertuples(index=False):
        ident = str(getattr(row, "ident", "") or "").upper()
        if not ident:
            continue
        lookup[ident] = {
            "ident": ident,
            "name": normalize_text(getattr(row, "name", "")) or ident,
            "lat": float(getattr(row, "latitude_deg")),
            "lon": float(getattr(row, "longitude_deg")),
            "source": normalize_text(getattr(row, "source", "")) or "",
        }
    return lookup


def airport_coord(ident: Any) -> dict[str, Any] | None:
    text = normalize_text(ident)
    if not text:
        return None
    return airport_coord_lookup().get(text.upper())


def _coord_dict_from_airport(ap: dict[str, Any] | None) -> dict[str, Any] | None:
    if not ap:
        return None
    return {"lat": float(ap["lat"]), "lon": float(ap["lon"]), "alt": None, "time": None}


def track_latlon_with_airport_extensions(
    row: pd.Series | dict[str, Any],
    points: list[dict[str, Any]],
    min_gap_km: float = 0.35,
) -> tuple[list[tuple[float, float]], list[dict[str, Any]]]:
    """Return a visual route extended to manually/automatically selected airports.

    KML from phones/trackers sometimes starts after departure or ends before
    arrival. We do not alter the stored GPS points or calculated GPS distance;
    only the map display gets a direct connector from/to known airports.
    """
    if not points:
        return [], []

    visual_points = [dict(p) for p in points]
    extensions: list[dict[str, Any]] = []

    dep_ap = airport_coord(row.get("departure") if hasattr(row, "get") else None)
    arr_ap = airport_coord(row.get("arrival") if hasattr(row, "get") else None)

    if dep_ap and visual_points:
        dep_pt = _coord_dict_from_airport(dep_ap)
        if dep_pt and haversine_km(dep_pt, visual_points[0]) > min_gap_km:
            visual_points.insert(0, dep_pt)
            extensions.append({"kind": "departure", "airport": dep_ap, "to": points[0]})

    if arr_ap and visual_points:
        arr_pt = _coord_dict_from_airport(arr_ap)
        if arr_pt and haversine_km(visual_points[-1], arr_pt) > min_gap_km:
            visual_points.append(arr_pt)
            extensions.append({"kind": "arrival", "airport": arr_ap, "from": points[-1]})

    latlon = [(float(p["lat"]), float(p["lon"])) for p in visual_points]
    return latlon, extensions


def map_center_from_tracks(tracks: pd.DataFrame) -> tuple[list[float], int]:
    coords: list[tuple[float, float]] = []
    for _, row in tracks.iterrows():
        try:
            points = downsample_points(json.loads(row["coordinates_json"]), max_points=300)
            latlon, _ = track_latlon_with_airport_extensions(row, points)
            stride = max(1, len(latlon) // 50 or 1)
            coords.extend(latlon[::stride])
        except Exception:
            continue
    if not coords:
        return [49.8, 15.5], 7
    min_lat = min(c[0] for c in coords); max_lat = max(c[0] for c in coords); min_lon = min(c[1] for c in coords); max_lon = max(c[1] for c in coords)
    center = [(min_lat + max_lat) / 2, (min_lon + max_lon) / 2]
    spread = max(max_lat - min_lat, max_lon - min_lon)
    zoom = 10 if spread < .25 else 8 if spread < 1 else 7 if spread < 4 else 6 if spread < 10 else 5
    return center, zoom


def make_map(
    tracks: pd.DataFrame,
    dark_mode: bool = True,
    line_weight: float = 4,
    line_opacity: float = 0.78,
    show_endpoints: bool = True,
    extend_to_airports: bool = True,
) -> folium.Map:
    center, zoom = map_center_from_tracks(tracks)
    tiles = "CartoDB dark_matter" if dark_mode else "OpenStreetMap"
    m = folium.Map(location=center, zoom_start=zoom, tiles=tiles, control_scale=True)
    if tracks.empty:
        return m
    for _, row in tracks.iterrows():
        try:
            points = downsample_points(json.loads(row["coordinates_json"]))
        except Exception:
            continue
        if len(points) < 2:
            continue
        gps_latlon = [(float(p["lat"]), float(p["lon"])) for p in points]
        latlon, extensions = track_latlon_with_airport_extensions(row, points) if extend_to_airports else (gps_latlon, [])
        evidence = str(row.get("evidence") or "").upper()
        color = "#38bdf8" if evidence == "ULL" else "#fbbf24"
        ext_note = "<br><span style='color:#94a3b8'>Mapa doplnila přímku k letišti.</span>" if extensions else ""
        popup = folium.Popup(f"""
            <b>{row.get('date') or ''} • {row.get('registration') or ''}</b><br>
            {row.get('departure') or ''}–{row.get('arrival') or ''}<br>
            {row.get('role') or ''} • {row.get('evidence') or ''}<br>
            GPS: {float(row.get('distance_km') or 0):.1f} km<br>
            Track: {row.get('file_name') or ''}{ext_note}
            """, max_width=360)
        folium.PolyLine(latlon, color=color, weight=line_weight, opacity=line_opacity, popup=popup).add_to(m)

        for ext in extensions:
            if ext["kind"] == "departure":
                seg = [(float(ext["airport"]["lat"]), float(ext["airport"]["lon"])), (float(ext["to"]["lat"]), float(ext["to"]["lon"]))]
                tooltip = f"Doplněno od letiště {ext['airport']['ident']} k prvnímu GPS bodu"
            else:
                seg = [(float(ext["from"]["lat"]), float(ext["from"]["lon"])), (float(ext["airport"]["lat"]), float(ext["airport"]["lon"]))]
                tooltip = f"Doplněno od posledního GPS bodu k letišti {ext['airport']['ident']}"
            folium.PolyLine(seg, color="#94a3b8", weight=max(1.2, line_weight - 0.6), opacity=0.72, dash_array="7,7", tooltip=tooltip).add_to(m)

        if show_endpoints:
            start = latlon[0] if latlon else gps_latlon[0]
            end = latlon[-1] if latlon else gps_latlon[-1]
            folium.CircleMarker(start, radius=4, color="#22c55e", fill=True, fill_opacity=.9, tooltip="Start / odlet").add_to(m)
            folium.CircleMarker(end, radius=4, color="#ef4444", fill=True, fill_opacity=.9, tooltip="End / přílet").add_to(m)
    folium.LayerControl().add_to(m)
    return m


def map_center_from_airport_coords(coords: list[tuple[float, float]]) -> tuple[list[float], int]:
    if not coords:
        return [49.8, 15.5], 7
    min_lat = min(c[0] for c in coords); max_lat = max(c[0] for c in coords); min_lon = min(c[1] for c in coords); max_lon = max(c[1] for c in coords)
    center = [(min_lat + max_lat) / 2, (min_lon + max_lon) / 2]
    spread = max(max_lat - min_lat, max_lon - min_lon)
    zoom = 11 if spread < .15 else 9 if spread < .6 else 8 if spread < 1.8 else 7 if spread < 5 else 6 if spread < 10 else 5
    return center, zoom


def make_route_overview_map(flights: pd.DataFrame, dark_mode: bool = True) -> folium.Map:
    lookup = airport_coord_lookup()
    coords: list[tuple[float, float]] = []
    routes: list[dict[str, Any]] = []
    visited: dict[str, dict[str, Any]] = {}

    for _, row in flights.iterrows():
        dep = normalize_text(row.get("departure"))
        arr = normalize_text(row.get("arrival"))
        if not dep or not arr:
            continue
        dep_ap = lookup.get(dep.upper())
        arr_ap = lookup.get(arr.upper())
        if not dep_ap or not arr_ap:
            continue
        dep_ll = (float(dep_ap["lat"]), float(dep_ap["lon"]))
        arr_ll = (float(arr_ap["lat"]), float(arr_ap["lon"]))
        coords.extend([dep_ll, arr_ll])
        for ap, kind in ((dep_ap, "dep"), (arr_ap, "arr")):
            ident = ap["ident"]
            if ident not in visited:
                visited[ident] = {**ap, "visits": 0, "departures": 0, "arrivals": 0, "first_date": "", "last_date": ""}
            visited[ident]["visits"] += 1
            if kind == "dep":
                visited[ident]["departures"] += 1
            else:
                visited[ident]["arrivals"] += 1
            date_txt = str(row.get("date") or "")
            if date_txt:
                if not visited[ident]["first_date"] or date_txt < visited[ident]["first_date"]:
                    visited[ident]["first_date"] = date_txt
                if not visited[ident]["last_date"] or date_txt > visited[ident]["last_date"]:
                    visited[ident]["last_date"] = date_txt
        routes.append({"row": row, "dep": dep_ap, "arr": arr_ap, "dep_ll": dep_ll, "arr_ll": arr_ll})

    center, zoom = map_center_from_airport_coords(coords)
    tiles = "CartoDB dark_matter" if dark_mode else "OpenStreetMap"
    m = folium.Map(location=center, zoom_start=zoom, tiles=tiles, control_scale=True)

    route_counts: dict[tuple[str, str], int] = {}
    for route in routes:
        row = route["row"]
        dep_id = route["dep"]["ident"]
        arr_id = route["arr"]["ident"]
        key = tuple(sorted([dep_id, arr_id]))
        route_counts[key] = route_counts.get(key, 0) + 1
        offset = min(route_counts[key] - 1, 8) * 0.0009
        dep_ll = (route["dep_ll"][0] + offset, route["dep_ll"][1] + offset)
        arr_ll = (route["arr_ll"][0] + offset, route["arr_ll"][1] + offset)
        evidence = str(row.get("evidence") or "").upper()
        color = "#38bdf8" if evidence == "ULL" else "#fbbf24"
        flight_id = int(row.get("id"))
        popup = folium.Popup(f"""
            <b>ID {flight_id} • {row.get('date') or ''}</b><br>
            {row.get('registration') or ''}<br>
            {dep_id}–{arr_id}<br>
            {row.get('off_block') or ''}–{row.get('on_block') or ''} • {row.get('role') or ''}
            """, max_width=320)
        tooltip = f"{dep_id}–{arr_id}"
        folium.PolyLine([dep_ll, arr_ll], color=color, weight=2.8, opacity=0.62, popup=popup, tooltip=tooltip).add_to(m)

    for ident, ap in visited.items():
        visits = int(ap.get("visits") or 0)
        radius = min(11, 4.5 + visits ** 0.5)
        tooltip = f"{ident} • {ap.get('name') or ''}"
        popup = folium.Popup(f"""
            <b>{ident}</b><br>
            {ap.get('name') or ''}<br>
            Návštěvy: {visits}<br>
            Odlety: {int(ap.get('departures') or 0)} • Přílety: {int(ap.get('arrivals') or 0)}<br>
            První: {ap.get('first_date') or '—'} • Poslední: {ap.get('last_date') or '—'}
            """, max_width=300)
        folium.CircleMarker((float(ap["lat"]), float(ap["lon"])), radius=radius, color="#22c55e", fill=True, fill_opacity=.92, tooltip=tooltip, popup=popup).add_to(m)

    folium.LayerControl().add_to(m)
    return m


def render_folium_readonly(m: folium.Map, *, height: int = 680, key: str | None = None) -> None:
    try:
        html = m.get_root().render()
        components.html(html, height=height, scrolling=False)
    except Exception:
        try:
            st_folium(m, height=height, use_container_width=True, key=key, returned_objects=[])
        except TypeError:
            st_folium(m, height=height, use_container_width=True, key=key)


def render_folium_navigable(m: folium.Map, *, height: int = 680, key: str | None = None) -> Any:
    try:
        return st_folium(
            m,
            height=height,
            use_container_width=True,
            key=key,
            returned_objects=["last_object_clicked", "last_object_clicked_tooltip", "last_object_clicked_popup"],
        )
    except TypeError:
        return st_folium(m, height=height, use_container_width=True, key=key)


def _stringify_map_event(value: Any) -> str:
    try:
        return json.dumps(value, ensure_ascii=False, default=str)
    except Exception:
        return str(value or "")


def handle_route_map_interaction(value: Any) -> None:
    text = _stringify_map_event(value)
    if not text:
        return

    route_match = re.search(r"([A-Z0-9]{3,5})\s*[–-]\s*([A-Z0-9]{3,5})", text)
    airport_match = re.search(r"\b([A-Z]{2}[A-Z0-9]{2,3}|[A-Z0-9]{3,5})\s*•", text)

    if route_match:
        dep = route_match.group(1).upper().strip()
        arr = route_match.group(2).upper().strip()
        if dep and arr and dep != arr:
            route = f"{dep}__{arr}"
            action = f"route:{route}"
            if st.session_state.get("_last_map_action") == action:
                return
            st.session_state["_last_map_action"] = action
            st.session_state["map_route"] = route
            st.session_state["map_mode_v035"] = "Orientační mapa letišť"
            st.session_state.pop("map_airport", None)
            st.rerun()

    if airport_match:
        ident = airport_match.group(1).upper().strip()
        if ident:
            action = f"airport:{ident}"
            if st.session_state.get("_last_map_action") == action:
                return
            st.session_state["_last_map_action"] = action
            st.session_state["map_airport"] = ident
            st.session_state["map_mode_v035"] = "Orientační mapa letišť"
            st.session_state.pop("map_route", None)
            st.rerun()


def _df_to_records_json(df: pd.DataFrame, columns: list[str]) -> str:
    """Stable compact JSON for cached map rendering."""
    if df.empty:
        return "[]"
    use_cols = [c for c in columns if c in df.columns]
    return df[use_cols].fillna("").to_json(orient="records", force_ascii=False, date_format="iso")


@st.cache_data(show_spinner=False, ttl=300)
def cached_track_map_html(records_json: str, dark_mode: bool) -> str:
    df = pd.read_json(BytesIO(records_json.encode("utf-8")), orient="records") if records_json and records_json != "[]" else pd.DataFrame()
    if df.empty:
        return ""
    m = make_map(df, dark_mode=dark_mode, line_weight=2, line_opacity=0.46, show_endpoints=False, extend_to_airports=True)
    return m.get_root().render()


@st.cache_data(show_spinner=False, ttl=300)
def cached_route_overview_map_html(records_json: str, dark_mode: bool) -> str:
    df = pd.read_json(BytesIO(records_json.encode("utf-8")), orient="records") if records_json and records_json != "[]" else pd.DataFrame()
    if df.empty:
        return ""
    m = make_route_overview_map(df, dark_mode=dark_mode)
    return m.get_root().render()


def render_map_html(html: str, *, height: int = 680) -> None:
    if not html:
        st.info("Mapa nemá data k zobrazení.")
        return
    components.html(html, height=height, scrolling=False)


def render_lazy_table(title: str, data: pd.DataFrame, *, height: int = 360, expanded: bool = False) -> None:
    """Keep heavy tables out of the main render path unless the user needs them."""
    with st.expander(title, expanded=expanded):
        if data.empty:
            st.info("Tabulka je prázdná.")
        else:
            st.dataframe(data, hide_index=True, use_container_width=True, height=height)


def render_track_profile(points: list[dict[str, Any]]) -> None:
    prof = profile_from_points(points)
    if prof.empty:
        st.info("Track nemá data pro profil.")
        return
    x = prof["time_local"] if prof["time_local"].notna().any() else prof["distance_km"]
    x_title = "Čas" if prof["time_local"].notna().any() else "Vzdálenost km"

    fig = go.Figure()
    fig.add_trace(
        go.Scatter(
            x=x,
            y=prof["alt_ft"],
            mode="lines",
            name="Altitude ft",
            line=dict(color="#38bdf8", width=2.4),
            hovertemplate="%{x}<br>Altitude: %{y:.0f} ft<extra></extra>",
        )
    )
    speed_values = pd.to_numeric(prof.get("speed_smooth"), errors="coerce") if "speed_smooth" in prof else pd.Series(dtype=float)
    has_speed = speed_values.notna().any()
    if has_speed:
        fig.add_trace(
            go.Scatter(
                x=x,
                y=speed_values,
                mode="lines",
                name="GPS speed km/h",
                yaxis="y2",
                line=dict(color="#f59e0b", width=2.2),
                hovertemplate="%{x}<br>Speed: %{y:.0f} km/h<extra></extra>",
            )
        )
    fig.update_layout(
        title="Profil letu",
        xaxis_title=x_title,
        yaxis=dict(title="Altitude ft", rangemode="tozero"),
        yaxis2=dict(title="GPS speed km/h", overlaying="y", side="right", rangemode="tozero"),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
    )
    st.plotly_chart(plotly_layout(fig), use_container_width=True)


# -----------------------------------------------------------------------------
# Pages
# -----------------------------------------------------------------------------

def page_dashboard(df: pd.DataFrame):
    st.markdown("## Dashboard")
    filtered = apply_filters(df, "dash")
    s = build_summary(filtered)
    c1, c2, c3, c4 = st.columns(4)
    with c1: metric_card("Celkový nálet", fmt_minutes(s["total"]), f"{s['flights']} letů")
    with c2: metric_card("PIC", fmt_minutes(s["pic"]), f"ULL {fmt_minutes(s['pic_ull'])} • EASA {fmt_minutes(s['pic_easa'])}")
    with c3: metric_card("DUAL / Safety", f"{fmt_minutes(s['dual'])} / {fmt_minutes(s['safety'])}", f"Starty {s['starts']}")
    with c4: metric_card("Náklady", fmt_money(s["cost"]), f"GPS {s['tracks']} tracků • {s['gps_km']:.0f} km")
    show_charts = st.toggle("Zobrazit grafy dashboardu", value=st.session_state.get("show_dashboard_charts", True), key="show_dashboard_charts")
    if not show_charts:
        return
    st.write("")
    chart_df = filtered.dropna(subset=["year"]).copy()
    if chart_df.empty:
        st.info("Žádná data pro grafy.")
        return
    yearly = chart_df.groupby("year", as_index=False).agg(Celkem=("block_hours", "sum"), Starty=("starts", "sum"), Naklady=("cost", "sum"))
    role_year = chart_df.pivot_table(index="year", columns="role", values="block_hours", aggfunc="sum", fill_value=0).reset_index()
    fig = px.bar(role_year, x="year", y=[c for c in ["PIC", "DUAL", "SAFETY PILOT", "INSTRUKTOR"] if c in role_year.columns], barmode="stack", title="Nálet podle roku a funkce")
    st.plotly_chart(plotly_layout(fig), use_container_width=True)
    left, right = st.columns(2)
    with left:
        by_aircraft = filtered.groupby("registration", as_index=False)["block_hours"].sum().sort_values("block_hours", ascending=False).head(12)
        fig2 = px.bar(by_aircraft, x="registration", y="block_hours", title="TOP letadla podle block time")
        st.plotly_chart(plotly_layout(fig2), use_container_width=True)
    with right:
        by_ev = filtered.groupby("evidence", as_index=False)["block_hours"].sum().sort_values("block_hours", ascending=False)
        fig3 = px.pie(by_ev, names="evidence", values="block_hours", title="ULL / EASA", hole=.45)
        st.plotly_chart(plotly_layout(fig3), use_container_width=True)


def flight_label(row: pd.Series | dict[str, Any]) -> str:
    return f"ID {int(row['id'])} • {row.get('date') or ''} • {row.get('registration') or ''} • {row.get('departure') or ''}-{row.get('arrival') or ''} • {row.get('off_block') or ''}-{row.get('on_block') or ''} • {row.get('role') or ''}"


def flight_display_df(df: pd.DataFrame) -> pd.DataFrame:
    cols = ["id","date","evidence","registration","aircraft_type","aircraft_class","departure","arrival","off_block","takeoff","landing","on_block","block_time","air_time","starts","commander","instructor","role","task","price_per_hour","cost_label","track_count","gps_km","note"]
    rename = {"id":"ID","date":"Datum","evidence":"Evidence","registration":"Imatrikulace","aircraft_type":"Typ","aircraft_class":"Třída","departure":"Odlet","arrival":"Přílet","off_block":"Off Block","takeoff":"Takeoff","landing":"Landing","on_block":"On Block","block_time":"Block","air_time":"Air","starts":"Starty","commander":"Velitel","instructor":"Instruktor","role":"Funkce","task":"Úloha","price_per_hour":"Kč/h","cost_label":"Cena","track_count":"GPS","gps_km":"GPS km","note":"Poznámka"}
    use = [c for c in cols if c in df.columns]
    return df[use].rename(columns=rename)


def flight_form(prefix: str, defaults: dict[str, Any], rates: pd.DataFrame, submit_label: str) -> dict[str, Any] | None:
    reg = str(defaults.get("registration") or "").upper()
    rate = lookup_latest_rate(rates, reg)
    default_price = defaults.get("price_per_hour") or rate.get("price_per_hour") or 0.0
    default_type = defaults.get("aircraft_type") or rate.get("aircraft_type") or ""
    with st.form(prefix):
        col1, col2, col3 = st.columns(3)
        with col1:
            flight_date = st.date_input("Datum", value=defaults.get("date") if isinstance(defaults.get("date"), date) else pd.to_datetime(defaults.get("date") or date.today()).date(), key=f"{prefix}_date")
            registration = st.text_input("Imatrikulace", value=reg, key=f"{prefix}_reg").upper()
            ev_def = defaults.get("evidence") or evidence_from_registration(reg)
            evidence = st.selectbox("Evidence", EVIDENCE_OPTIONS, index=EVIDENCE_OPTIONS.index(ev_def) if ev_def in EVIDENCE_OPTIONS else 0, key=f"{prefix}_ev")
            aircraft_type = st.text_input("Typ", value=str(default_type or ""), key=f"{prefix}_type")
            cls_def = defaults.get("aircraft_class") or default_class_for(evidence)
            aircraft_class = st.selectbox("Třída", CLASS_OPTIONS, index=CLASS_OPTIONS.index(cls_def) if cls_def in CLASS_OPTIONS else 0, key=f"{prefix}_class")
        with col2:
            departure = st.text_input("Odlet", value=str(defaults.get("departure") or ""), key=f"{prefix}_dep").upper()
            arrival = st.text_input("Přílet", value=str(defaults.get("arrival") or ""), key=f"{prefix}_arr").upper()
            off_block = st.text_input("Off Block", value=str(defaults.get("off_block") or ""), key=f"{prefix}_off")
            takeoff = st.text_input("Takeoff", value=str(defaults.get("takeoff") or ""), key=f"{prefix}_to")
            landing = st.text_input("Landing", value=str(defaults.get("landing") or ""), key=f"{prefix}_ldg")
            on_block = st.text_input("On Block", value=str(defaults.get("on_block") or ""), key=f"{prefix}_on")
        with col3:
            starts = st.number_input("Starty", min_value=0, step=1, value=int(defaults.get("starts") or 1), key=f"{prefix}_starts")
            commander = st.text_input("Velitel", value=str(defaults.get("commander") or "Točík Filip"), key=f"{prefix}_cmd")
            instructor = st.text_input("Instruktor", value=str(defaults.get("instructor") or ""), key=f"{prefix}_instr")
            role_def = defaults.get("role") or "PIC"
            role = st.selectbox("Funkce", ROLE_OPTIONS, index=ROLE_OPTIONS.index(role_def) if role_def in ROLE_OPTIONS else 0, key=f"{prefix}_role")
            price = st.number_input("Cena Kč/h", min_value=0.0, step=50.0, value=float(default_price or 0), key=f"{prefix}_price")
            task = st.text_input("Úloha", value=str(defaults.get("task") or ""), key=f"{prefix}_task")
            note = st.text_input("Poznámka", value=str(defaults.get("note") or ""), key=f"{prefix}_note")
        block = minutes_diff(off_block, on_block); air = minutes_diff(takeoff, landing)
        c1, c2, c3 = st.columns(3)
        with c1: metric_card("Block Time", fmt_minutes(block), "")
        with c2: metric_card("Air Time", fmt_minutes(air), "")
        with c3: metric_card("Cena letu", fmt_money((block or 0)/60*price), "")
        submitted = st.form_submit_button(submit_label, type="primary", use_container_width=True)
    if submitted:
        return {"date": flight_date, "evidence": evidence, "registration": registration, "aircraft_type": aircraft_type, "aircraft_class": aircraft_class, "departure": departure, "arrival": arrival, "off_block": off_block, "takeoff": takeoff, "landing": landing, "on_block": on_block, "starts": int(starts), "commander": commander, "instructor": instructor, "role": role, "task": task, "price_per_hour": price, "note": note}
    return None



@st.dialog("Detail letu", width="large", dismissible=True, on_dismiss=clear_open_flight_dialog)
def flight_detail_dialog(selected_id: int, row_data: dict[str, Any], rates: pd.DataFrame, dark_mode: bool) -> None:
    row = pd.Series(row_data)
    route_text = f"{_safe_text(row_data.get('departure')) or '—'} → {_safe_text(row_data.get('arrival')) or '—'}"
    st.markdown(
        f"""
        <div class="flight-detail-hero">
            <div class="flight-detail-route">{route_text}</div>
            <div class="flight-detail-meta">
                <span>ID {int(selected_id)}</span>
                <span>{_safe_text(row_data.get('date')) or 'bez data'}</span>
                <span>{_safe_text(row_data.get('registration')) or 'bez imatrikulace'}</span>
                <span>{_safe_text(row_data.get('role')) or 'bez funkce'}</span>
                <span>{_safe_text(row_data.get('evidence')) or 'bez evidence'}</span>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    detail_section = st.radio(
        "Sekce detailu",
        ["Přehled", "Editace", "Track", "Smazání"],
        horizontal=True,
        key=f"detail_section_{selected_id}",
        label_visibility="collapsed",
    )
    if detail_section == "Přehled":
        block_text = _safe_text(row.get("block_time")) or "—"
        air_text = _safe_text(row.get("air_time")) or "—"
        price_text = _safe_text(row.get("cost_label")) or "—"
        gps_count = int(_safe_float(row.get("track_count"), 0))
        gps_km = _safe_float(row.get("gps_km"), 0)
        st.markdown(
            f"""
            <div class="detail-grid">
              <div class="detail-card">
                <div class="detail-card-label">Čas letu</div>
                <div class="detail-card-value">{block_text}</div>
                <div class="detail-card-sub">Air {air_text}</div>
              </div>
              <div class="detail-card">
                <div class="detail-card-label">Letadlo</div>
                <div class="detail-card-value">{_safe_text(row.get('registration')) or '—'}</div>
                <div class="detail-card-sub">{_safe_text(row.get('aircraft_type')) or '—'} · {_safe_text(row.get('aircraft_class')) or '—'}</div>
              </div>
              <div class="detail-card">
                <div class="detail-card-label">Cena</div>
                <div class="detail-card-value">{price_text}</div>
                <div class="detail-card-sub">{_safe_text(row.get('price_per_hour')) or '—'} Kč/h</div>
              </div>
              <div class="detail-card">
                <div class="detail-card-label">GPS</div>
                <div class="detail-card-value">{gps_count}</div>
                <div class="detail-card-sub">{gps_km:.1f} km</div>
              </div>
            </div>
            <div class="detail-split">
              <div class="detail-kv">
                <div class="detail-kv-title">Let</div>
                <div class="detail-kv-row"><span>Datum</span><span>{_safe_text(row.get('date')) or '—'}</span></div>
                <div class="detail-kv-row"><span>Trasa</span><span>{_safe_text(row.get('departure')) or '—'} → {_safe_text(row.get('arrival')) or '—'}</span></div>
                <div class="detail-kv-row"><span>Evidence</span><span>{_safe_text(row.get('evidence')) or '—'}</span></div>
                <div class="detail-kv-row"><span>Funkce</span><span>{_safe_text(row.get('role')) or '—'}</span></div>
                <div class="detail-kv-row"><span>Starty</span><span>{_safe_text(row.get('starts')) or '—'}</span></div>
              </div>
              <div class="detail-kv">
                <div class="detail-kv-title">Časy a posádka</div>
                <div class="detail-kv-row"><span>Block</span><span>{_safe_text(row.get('off_block')) or '—'} – {_safe_text(row.get('on_block')) or '—'}</span></div>
                <div class="detail-kv-row"><span>Air</span><span>{_safe_text(row.get('takeoff')) or '—'} – {_safe_text(row.get('landing')) or '—'}</span></div>
                <div class="detail-kv-row"><span>Velitel</span><span>{_safe_text(row.get('commander')) or '—'}</span></div>
                <div class="detail-kv-row"><span>Instruktor</span><span>{_safe_text(row.get('instructor')) or '—'}</span></div>
                <div class="detail-kv-row"><span>Úloha</span><span>{_safe_text(row.get('task')) or '—'}</span></div>
              </div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        note_text = _safe_text(row.get("note"))
        if note_text:
            st.markdown(f'<div class="detail-kv"><div class="detail-kv-title">Poznámka</div>{note_text}</div>', unsafe_allow_html=True)
    elif detail_section == "Editace":
        if not is_admin():
            st.info("Pouze admin.")
        else:
            saved = flight_form(f"edit_flight_{selected_id}", row.to_dict(), rates, "Uložit změny")
            if saved is not None:
                update_flight(int(selected_id), saved)
                st.success("Změny uloženy.")
                clear_open_flight_dialog()
                st.rerun()
    elif detail_section == "Track":
        flight_tracks = read_tracks_for_flight(int(selected_id))
        if not flight_tracks.empty:
            joined = read_tracks_joined()
            render_folium_readonly(make_map(joined[joined["flight_id"].eq(int(selected_id))], dark_mode), height=440, key=f"track_map_existing_{selected_id}_{len(flight_tracks)}")
            first_points = json.loads(flight_tracks.iloc[0]["coordinates_json"])
            render_track_profile(first_points)
            show = flight_tracks[["id","file_name","point_count","distance_km","start_utc","end_utc","max_alt_m"]].rename(columns={"id":"Track ID","file_name":"Soubor","point_count":"Body","distance_km":"Km","start_utc":"Start UTC","end_utc":"End UTC","max_alt_m":"Max alt m"})
            st.dataframe(show, hide_index=True, use_container_width=True)
            del_id = st.selectbox("Smazat track", show["Track ID"].tolist(), format_func=lambda x: f"Track ID {x}")
            if st.button("Smazat vybraný track", type="secondary", disabled=not is_admin(), use_container_width=True):
                if require_admin():
                    delete_track(int(del_id)); st.success("Track smazán."); st.rerun()
        else:
            st.info("K letu zatím není připojený track.")
        uploaded = st.file_uploader("Přidat / nahradit KML track", type=["kml"], key=f"attach_track_{selected_id}")
        if uploaded is not None:
            try:
                points = parse_kml_bytes(uploaded.read())
                if len(points) >= 2:
                    preview = pd.DataFrame([{"id": -1,"flight_id": selected_id,"coordinates_json": json.dumps(points),"file_name": uploaded.name,"distance_km": track_stats(points)["distance_km"],"date": row.get("date"),"registration": row.get("registration"),"departure": row.get("departure"),"arrival": row.get("arrival"),"role": row.get("role"),"evidence": row.get("evidence")}])
                    render_folium_readonly(make_map(preview, dark_mode), height=360, key=f"track_map_preview_{selected_id}_{uploaded.name}_{len(points)}")
                    replace = st.checkbox("Nahradit existující tracky u tohoto letu", value=True, key=f"replace_track_{selected_id}_{uploaded.name}")
                    if st.button("Uložit track k letu", type="primary", disabled=not is_admin(), use_container_width=True):
                        if require_admin():
                            save_track(int(selected_id), uploaded.name, points, replace_existing=replace)
                            st.success("Track uložen."); st.rerun()
                else:
                    st.error("V KML nejsou použitelné body.")
            except Exception as exc:
                st.error(f"KML / náhled se nepodařilo zpracovat: {exc}")
    elif detail_section == "Smazání":
        if not is_admin():
            st.info("Pouze admin.")
        else:
            st.markdown(
                f"""
                <div class="danger-box">
                    <div class="danger-title">Trvalé smazání letu</div>
                    <div class="danger-text">Tato akce smaže let ID {selected_id} ze zápisníku, včetně všech připojených KML tracků a GPS bodů. Po uložení se změna automaticky zazálohuje na GitHub, pokud je záloha zapnutá.</div>
                </div>
                """,
                unsafe_allow_html=True,
            )
            confirm = st.text_input(
                f"Pro potvrzení napiš ID letu: {selected_id}",
                value="",
                key=f"delete_flight_confirm_{selected_id}",
            )
            c_del, c_cancel = st.columns([1, 1])
            with c_del:
                if st.button("Trvale smazat let", type="primary", use_container_width=True, key=f"delete_flight_btn_{selected_id}"):
                    if confirm.strip() == str(selected_id):
                        delete_flight(int(selected_id))
                        st.success(f"Let ID {selected_id} byl smazán.")
                        clear_open_flight_dialog()
                        st.rerun()
                    else:
                        st.error("Potvrzení nesouhlasí. Napiš přesné ID letu.")
            with c_cancel:
                if st.button("Nemazat", use_container_width=True, key=f"delete_flight_cancel_{selected_id}"):
                    clear_open_flight_dialog()
                    st.rerun()
    if st.button("Zavřít detail", use_container_width=True):
        clear_open_flight_dialog()
        st.rerun()



def _is_blank(value: Any) -> bool:
    """Return True for None, NaN/NaT/pd.NA and textual empty markers.

    User-added flights may have optional fields blank. The list view must never crash
    on missing aircraft type/class, route, task, price, etc.
    """
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except Exception:
        pass
    text = str(value).strip()
    return text == "" or text.lower() in {"none", "nan", "nat", "<na>"}


def _clean_text(value: Any) -> str:
    return "" if _is_blank(value) else str(value).strip()


def _safe_text(value: Any) -> str:
    text = _clean_text(value)
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _join_nonblank(values: list[Any] | tuple[Any, ...], sep: str = " • ") -> str:
    return sep.join(_clean_text(v) for v in values if not _is_blank(v))


def _range_text(start: Any, end: Any, sep: str = "–") -> str:
    left = _clean_text(start)
    right = _clean_text(end)
    if left and right:
        return f"{left}{sep}{right}"
    return left or right


def _safe_float(value: Any, default: float = 0.0) -> float:
    if _is_blank(value):
        return default
    try:
        number = float(value)
        if math.isnan(number) or math.isinf(number):
            return default
        return number
    except Exception:
        return default


def _safe_int(value: Any, default: int = 0) -> int:
    try:
        return int(round(_safe_float(value, float(default))))
    except Exception:
        return default


def _price_rate_label(value: Any) -> str:
    if _is_blank(value):
        return ""
    return f"{_safe_float(value):.0f} Kč/h"


def _cell(main: Any, sub: Any = "") -> str:
    main_txt = _safe_text(main)
    sub_txt = _safe_text(sub)
    if sub_txt:
        return f'<div class="flight-cell"><div class="flight-cell-main">{main_txt}</div><div class="flight-cell-sub">{sub_txt}</div></div>'
    return f'<div class="flight-cell"><div class="flight-cell-main">{main_txt}</div></div>'


def render_flight_list(table_df: pd.DataFrame, rates: pd.DataFrame, dark_mode: bool) -> None:
    """Compact paginated flight list with one real Detail button per visible row.

    This avoids unreliable table selection and avoids opening browser links. Only the
    visible page gets buttons, so it stays responsive on Streamlit Cloud.
    """
    if table_df.empty:
        st.info("Filtr nevrátil žádné lety.")
        return

    display_df = flight_display_df(table_df).copy()

    controls = st.columns([1.0, 2.4, 1.0, 1.0])
    with controls[0]:
        page_size_choice = st.selectbox("Řádků", [25, 50, 100, "Vše"], index=0, key="flight_page_size_v14")
    with controls[1]:
        quick_filter = st.text_input("Rychlé hledání", value="", placeholder="registrace, letiště, typ, funkce…", key="flight_table_quick_filter_v14")

    if quick_filter.strip():
        q = quick_filter.strip().lower()
        mask = display_df.astype(str).apply(lambda col: col.str.lower().str.contains(q, na=False)).any(axis=1)
        shown_display = display_df[mask].copy()
        shown_table = table_df.loc[shown_display.index].copy()
    else:
        shown_display = display_df.copy()
        shown_table = table_df.copy()

    shown_display = shown_display.reset_index(drop=True)
    shown_table = shown_table.reset_index(drop=True)

    total_rows = len(shown_table)
    show_all_rows = page_size_choice == "Vše"
    page_size = total_rows if show_all_rows else int(page_size_choice)
    page_size = max(1, page_size)
    page_count = max(1, math.ceil(total_rows / page_size))
    with controls[2]:
        if show_all_rows:
            page = 1
            st.text_input("Stránka", value="Vše", disabled=True, key="flight_page_all_v14")
        else:
            page = st.number_input("Stránka", min_value=1, max_value=page_count, value=min(int(st.session_state.get("flight_page_v14", page_count)), page_count), step=1, key="flight_page_v14")
    with controls[3]:
        st.markdown(f'<div class="flight-page-info">{total_rows} letů • {page_count} stran</div>', unsafe_allow_html=True)

    start = 0 if show_all_rows else (int(page) - 1) * page_size
    end = total_rows if show_all_rows else start + page_size
    page_rows = shown_table.iloc[start:end].copy()

    widths = [0.72, 0.48, 0.88, 0.62, 1.25, 1.02, 1.05, 0.82, 0.52, 0.92, 1.22, 0.92, 0.78, 0.50]
    headers = ["Detail", "ID", "Datum", "Ev.", "Letadlo", "Trasa", "Časy", "Block", "St.", "Funkce", "Velitel", "Úloha", "Cena", "GPS"]
    hcols = st.columns(widths, gap="small", vertical_alignment="top")
    for col, header in zip(hcols, headers):
        col.markdown(f'<div class="flight-list-head">{header}</div>', unsafe_allow_html=True)

    for _, row in page_rows.iterrows():
        flight_id = int(row.get("id"))
        cols = st.columns(widths, gap="small", vertical_alignment="top")
        with cols[0]:
            if st.button("Detail", key=f"flight_detail_btn_{flight_id}", use_container_width=True):
                st.session_state["open_flight_dialog_id"] = flight_id
                st.session_state["selected_flight_id"] = flight_id
                st.session_state.pop("dismissed_flight_id", None)
        cols[1].markdown(_cell(flight_id), unsafe_allow_html=True)
        cols[2].markdown(_cell(row.get("date")), unsafe_allow_html=True)
        cols[3].markdown(_cell(row.get("evidence")), unsafe_allow_html=True)
        aircraft_sub = _join_nonblank([row.get("aircraft_type"), row.get("aircraft_class")])
        cols[4].markdown(_cell(row.get("registration"), aircraft_sub), unsafe_allow_html=True)
        cols[5].markdown(_cell(_range_text(row.get("departure"), row.get("arrival"))), unsafe_allow_html=True)
        time_main = _range_text(row.get("off_block"), row.get("on_block"))
        air_range = _range_text(row.get("takeoff"), row.get("landing"))
        time_sub = f"Air {air_range}" if air_range else ""
        cols[6].markdown(_cell(time_main, time_sub), unsafe_allow_html=True)
        cols[7].markdown(_cell(row.get("block_time"), f"Air {row.get('air_time') or ''}"), unsafe_allow_html=True)
        cols[8].markdown(_cell(_safe_int(row.get("starts"))), unsafe_allow_html=True)
        cols[9].markdown(_cell(row.get("role")), unsafe_allow_html=True)
        cols[10].markdown(_cell(row.get("commander"), row.get("instructor") if not _is_blank(row.get("instructor")) else ""), unsafe_allow_html=True)
        cols[11].markdown(_cell(row.get("task")), unsafe_allow_html=True)
        cols[12].markdown(_cell(row.get("cost_label"), _price_rate_label(row.get("price_per_hour"))), unsafe_allow_html=True)
        gps_km = _safe_float(row.get("gps_km"))
        cols[13].markdown(_cell(_safe_int(row.get("track_count")), f"{gps_km:.0f} km"), unsafe_allow_html=True)
        st.markdown('<div class="flight-row-sep"></div>', unsafe_allow_html=True)

    open_id = st.session_state.get("open_flight_dialog_id")
    valid_ids = set(table_df["id"].astype(int).tolist())
    if open_id is not None and int(open_id) in valid_ids:
        dialog_row = table_df[table_df["id"].astype(int).eq(int(open_id))].iloc[0]
        flight_detail_dialog(int(open_id), dialog_row.to_dict(), rates, dark_mode)


def page_logbook(df: pd.DataFrame, rates: pd.DataFrame, dark_mode: bool):
    st.markdown("## Lety")
    filtered = apply_filters(df, "logbook")
    s = build_summary(filtered)
    c1, c2, c3, c4 = st.columns(4)
    with c1: metric_card("Zobrazeno", str(s["flights"]), "letů")
    with c2: metric_card("Celkem", fmt_minutes(s["total"]), "block time")
    with c3: metric_card("PIC", fmt_minutes(s["pic"]), "z filtrovaných letů")
    with c4: metric_card("GPS", str(s["tracks"]), f"{s['gps_km']:.0f} km")

    if filtered.empty:
        st.info("Filtr nevrátil žádné lety.")
        return

    table_df = filtered.sort_values(["date_dt", "off_block", "id"], na_position="last").reset_index(drop=True)
    render_flight_list(table_df, rates, dark_mode)

def page_new_flight(rates: pd.DataFrame, dark_mode: bool):
    st.markdown("## Nový let")
    if not is_admin():
        st.info("Pouze admin.")
        return
    tab_track, tab_manual = st.tabs(["Z tracku", "Ručně"])
    with tab_track:
        uploaded = st.file_uploader("KML track", type=["kml"], key="new_track_kml")
        if uploaded is not None:
            raw = uploaded.read()
            try:
                points = parse_kml_bytes(raw)
            except Exception as exc:
                st.error(f"KML / náhled se nepodařilo zpracovat: {exc}")
                points = []
            if len(points) >= 2:
                defaults = infer_from_track(points, uploaded.name, rates)
                stats = defaults.pop("stats")
                detect_idx = defaults.pop("detect_idx", {})
                has_clock = defaults.pop("has_clock", False)
                c1, c2, c3, c4 = st.columns(4)
                with c1: metric_card("Body", str(stats["point_count"]), uploaded.name)
                with c2: metric_card("GPS délka", f"{stats['distance_km']:.1f} km", "")
                with c3: metric_card("Časy", f"{defaults.get('takeoff') or '—'}–{defaults.get('landing') or '—'}", "Takeoff / landing")
                with c4: metric_card("Block", f"{defaults.get('off_block') or '—'}–{defaults.get('on_block') or '—'}", "automaticky ±5 min")
                if not has_clock:
                    st.warning("KML neobsahuje časové značky u bodů. Časy doplň ručně.")
                preview_df = pd.DataFrame([{"id": -1,"flight_id": -1,"coordinates_json": json.dumps(points),"file_name": uploaded.name,"distance_km": stats["distance_km"],"date": defaults.get("date"),"registration": defaults.get("registration"),"departure": defaults.get("departure"),"arrival": defaults.get("arrival"),"role": defaults.get("role"),"evidence": defaults.get("evidence")}])
                render_folium_readonly(make_map(preview_df, dark_mode), height=420, key=f"new_flight_preview_map_{uploaded.name}_{len(points)}")
                with st.expander("Profil tracku", expanded=True):
                    render_track_profile(points)
                saved = flight_form("new_from_track", defaults, rates, "Uložit nový let včetně tracku")
                if saved is not None:
                    flight_id = create_flight(saved, auto_backup=False)
                    save_track(flight_id, uploaded.name, points, replace_existing=True)
                    st.success(f"Let uložen jako ID {flight_id}.")
                    st.rerun()
            elif uploaded is not None:
                st.error("V KML nejsou použitelné body trasy.")
    with tab_manual:
        defaults = {"date": date.today(), "evidence": "ULL", "aircraft_class": "ULL", "starts": 1, "commander": "Točík Filip", "role": "PIC"}
        saved = flight_form("new_manual", defaults, rates, "Přidat let")
        if saved is not None:
            flight_id = create_flight(saved)
            st.success(f"Let uložen jako ID {flight_id}.")
            st.rerun()




def map_navigation_options(df: pd.DataFrame) -> tuple[list[str], list[tuple[str, str]], int]:
    lookup = airport_coord_lookup()
    airports: set[str] = set()
    routes: dict[tuple[str, str], int] = {}
    known_routes = 0
    if df.empty:
        return [], [], 0
    for _, row in df.iterrows():
        dep = normalize_text(row.get("departure"))
        arr = normalize_text(row.get("arrival"))
        if not dep or not arr:
            continue
        dep = dep.upper()
        arr = arr.upper()
        if dep not in lookup or arr not in lookup:
            continue
        known_routes += 1
        airports.add(dep)
        airports.add(arr)
        key = tuple(sorted([dep, arr]))
        routes[key] = routes.get(key, 0) + 1
    airport_options = sorted(airports)
    route_options = []
    for (dep, arr), count in sorted(routes.items(), key=lambda item: (-item[1], item[0])):
        value = f"{dep}__{arr}"
        label = f"{dep}–{arr} ({count})"
        route_options.append((value, label))
    return airport_options, route_options, known_routes


def render_map_navigation_controls(filtered: pd.DataFrame) -> None:
    airports, routes, _ = map_navigation_options(filtered)
    current_airport = normalize_text(st.session_state.get("map_airport"))
    current_route = normalize_text(st.session_state.get("map_route"))
    if current_airport and current_airport not in airports:
        airports = [current_airport] + airports
    route_values_existing = [value for value, _ in routes]
    if current_route and current_route not in route_values_existing:
        routes = [(current_route, current_route.replace("__", "–"))] + routes
    if not airports and not routes:
        return
    c1, c2 = st.columns(2)
    with c1:
        airport_choice = st.selectbox("Letiště", [""] + airports, index=0, key="map_airport_picker_v0356", format_func=lambda v: v or "—")
    route_values = [""] + [value for value, _ in routes]
    route_labels = {value: label for value, label in routes}
    with c2:
        route_choice = st.selectbox("Trasa", route_values, index=0, key="map_route_picker_v0356", format_func=lambda v: route_labels.get(v, "—"))
    if airport_choice:
        st.session_state["map_airport"] = airport_choice
        st.session_state.pop("map_route", None)
    elif route_choice:
        st.session_state["map_route"] = route_choice
        st.session_state.pop("map_airport", None)

def _airport_selection(df: pd.DataFrame, ident: str) -> pd.DataFrame:
    if df.empty or not ident:
        return pd.DataFrame()
    ident = str(ident).upper().strip()
    dep = df.get("departure", pd.Series(dtype=str)).fillna("").astype(str).str.upper()
    arr = df.get("arrival", pd.Series(dtype=str)).fillna("").astype(str).str.upper()
    return df[dep.eq(ident) | arr.eq(ident)].copy()


def _parse_route_selection(value: str | None) -> tuple[str, str] | None:
    if not value or "__" not in str(value):
        return None
    dep, arr = str(value).upper().split("__", 1)
    dep = dep.strip()
    arr = arr.strip()
    if not dep or not arr:
        return None
    return dep, arr


def _route_selection(df: pd.DataFrame, dep: str, arr: str) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()
    dep = dep.upper().strip()
    arr = arr.upper().strip()
    d = df.get("departure", pd.Series(dtype=str)).fillna("").astype(str).str.upper()
    a = df.get("arrival", pd.Series(dtype=str)).fillna("").astype(str).str.upper()
    return df[(d.eq(dep) & a.eq(arr)) | (d.eq(arr) & a.eq(dep))].copy()


def _clear_map_selection() -> None:
    st.session_state.pop("map_airport", None)
    st.session_state.pop("map_route", None)
    st.session_state.pop("_last_map_action", None)
    try:
        for key in ("map_airport", "map_route"):
            if key in st.query_params:
                del st.query_params[key]
    except Exception:
        pass


def render_map_selection_panel(selection_df: pd.DataFrame, title: str, rates: pd.DataFrame, dark_mode: bool) -> None:
    if selection_df.empty:
        st.markdown(f'<div class="map-selection-panel"><div class="map-selection-title">{title}</div></div>', unsafe_allow_html=True)
        return
    work = selection_df.sort_values(["date_dt", "off_block", "id"], ascending=[False, False, False], na_position="last").reset_index(drop=True)
    total_minutes = int(work.get("block_minutes", pd.Series(dtype=float)).fillna(0).sum()) if "block_minutes" in work else 0
    tracks = int(work.get("track_count", pd.Series(dtype=float)).fillna(0).sum()) if "track_count" in work else 0
    gps = float(work.get("gps_km", pd.Series(dtype=float)).fillna(0).sum()) if "gps_km" in work else 0.0
    st.markdown(
        f"""<div class="map-selection-panel">
            <div class="map-selection-title">{title}</div>
            <div class="map-selection-meta">
                <span>{len(work)} letů</span><span>{fmt_minutes(total_minutes)}</span><span>{tracks} tracků</span><span>{gps:.0f} km GPS</span>
            </div>
        </div>""",
        unsafe_allow_html=True,
    )
    top = st.columns([.62,.58,.55,.86,1.1,1.05,.86,.72,.75], gap="small")
    headers = ["Detail", "Track", "ID", "Datum", "Letadlo", "Trasa", "Časy", "Block", "Role"]
    for c, h in zip(top, headers):
        c.markdown(f'<div class="map-mini-head">{h}</div>', unsafe_allow_html=True)
    limit = min(24, len(work))
    for _, row in work.head(limit).iterrows():
        flight_id = int(row.get("id"))
        cols = st.columns([.62,.58,.55,.86,1.1,1.05,.86,.72,.75], gap="small", vertical_alignment="top")
        with cols[0]:
            if st.button("Detail", key=f"map_selection_detail_{flight_id}", use_container_width=True):
                st.session_state[f"detail_section_{flight_id}"] = "Přehled"
                st.session_state["open_flight_dialog_id"] = flight_id
                st.session_state["selected_flight_id"] = flight_id
                st.session_state.pop("dismissed_flight_id", None)
                st.rerun()
        with cols[1]:
            track_count = int(_safe_float(row.get("track_count"), 0))
            if st.button("Track", key=f"map_selection_track_{flight_id}", disabled=track_count <= 0, use_container_width=True):
                st.session_state[f"detail_section_{flight_id}"] = "Track"
                st.session_state["open_flight_dialog_id"] = flight_id
                st.session_state["selected_flight_id"] = flight_id
                st.session_state.pop("dismissed_flight_id", None)
                st.rerun()
        cols[2].markdown(f'<div class="map-mini-cell">{flight_id}</div>', unsafe_allow_html=True)
        cols[3].markdown(f'<div class="map-mini-cell">{_safe_text(row.get("date"))}</div>', unsafe_allow_html=True)
        cols[4].markdown(f'<div class="map-mini-cell">{_safe_text(row.get("registration"))}<div class="map-mini-sub">{_safe_text(row.get("aircraft_type"))}</div></div>', unsafe_allow_html=True)
        cols[5].markdown(f'<div class="map-mini-cell">{_range_text(row.get("departure"), row.get("arrival"))}<div class="map-mini-sub">{_safe_text(row.get("evidence"))}</div></div>', unsafe_allow_html=True)
        cols[6].markdown(f'<div class="map-mini-cell">{_range_text(row.get("off_block"), row.get("on_block"))}</div>', unsafe_allow_html=True)
        cols[7].markdown(f'<div class="map-mini-cell">{_safe_text(row.get("block_time"))}</div>', unsafe_allow_html=True)
        cols[8].markdown(f'<div class="map-mini-cell">{_safe_text(row.get("role"))}</div>', unsafe_allow_html=True)
    if len(work) > limit:
        render_lazy_table(
            "Všechny vybrané lety",
            work[["id","date","registration","departure","arrival","role","evidence","block_time","track_count","gps_km"]]
            .rename(columns={"id":"ID","date":"Datum","registration":"Imatrikulace","departure":"Odlet","arrival":"Přílet","role":"Funkce","evidence":"Evidence","block_time":"Block","track_count":"Tracky","gps_km":"GPS km"}),
            height=360,
        )
    open_id = st.session_state.get("open_flight_dialog_id")
    if open_id is not None and int(open_id) in set(work["id"].astype(int).tolist()):
        dialog_row = work[work["id"].astype(int).eq(int(open_id))].iloc[0]
        flight_detail_dialog(int(open_id), dialog_row.to_dict(), rates, dark_mode)


def render_map_selection(filtered: pd.DataFrame, rates: pd.DataFrame, dark_mode: bool) -> None:
    airport = normalize_text(st.session_state.get("map_airport"))
    route = _parse_route_selection(st.session_state.get("map_route"))
    if not airport and not route:
        return
    if airport:
        ident = airport.upper()
        selected = _airport_selection(filtered, ident)
        render_map_selection_panel(selected, f"Letiště {ident}", rates, dark_mode)
    elif route:
        dep, arr = route
        selected = _route_selection(filtered, dep, arr)
        render_map_selection_panel(selected, f"Trasa {dep}–{arr}", rates, dark_mode)

def page_maps(flights: pd.DataFrame, rates: pd.DataFrame, dark_mode: bool):
    st.markdown("## Mapa letů")
    filtered = apply_filters(flights, "map")
    st.markdown('<div class="map-mode-row">', unsafe_allow_html=True)
    map_mode = st.radio(
        "Typ mapy",
        ["Orientační mapa letišť", "GPS tracky"],
        horizontal=True,
        label_visibility="collapsed",
        key="map_mode_v035",
    )
    st.markdown('</div>', unsafe_allow_html=True)

    base_track_count = int(filtered.get("track_count", pd.Series(dtype=float)).fillna(0).sum()) if not filtered.empty else 0
    base_gps_km = float(filtered.get("gps_km", pd.Series(dtype=float)).fillna(0).sum()) if not filtered.empty else 0.0
    known_routes: int | None = None
    if map_mode == "Orientační mapa letišť":
        _, _, known_routes = map_navigation_options(filtered)
        render_map_navigation_controls(filtered)
        if st.session_state.get("map_airport") or st.session_state.get("map_route"):
            _, clear_col = st.columns([1, .16])
            with clear_col:
                if st.button("Zrušit", key="map_selection_clear", use_container_width=True):
                    _clear_map_selection()
                    st.rerun()

    c1, c2, c3, c4 = st.columns(4)
    with c1: metric_card("Letů ve filtru", str(len(filtered)), "")
    with c2: metric_card("Tracky", str(base_track_count), "")
    with c3: metric_card("GPS vzdálenost", f"{base_gps_km:.1f} km", "")
    with c4: metric_card("Direct trasy", str(known_routes) if known_routes is not None else "—", "")

    if map_mode == "GPS tracky":
        tracks = read_tracks_joined()
        if not tracks.empty:
            tracks = tracks[tracks["flight_id"].isin(filtered["id"].tolist())].copy()
        if tracks.empty:
            st.info("Pro aktuální filtr není dostupný žádný KML track.")
        else:
            records_json = _df_to_records_json(
                tracks,
                ["flight_id", "id", "date", "registration", "departure", "arrival", "role", "evidence", "file_name", "point_count", "distance_km", "coordinates_json"],
            )
            render_map_html(cached_track_map_html(records_json, bool(dark_mode)), height=680)
            render_lazy_table(
                "Tabulka GPS tracků",
                tracks[["date","registration","departure","arrival","role","evidence","file_name","point_count","distance_km"]].rename(columns={"date":"Datum","registration":"Imatrikulace","departure":"Odlet","arrival":"Přílet","role":"Funkce","evidence":"Evidence","file_name":"Soubor","point_count":"Body","distance_km":"Km"}),
                height=320,
            )
    else:
        if filtered.empty or not known_routes:
            st.info("Pro aktuální filtr nejsou známé souřadnice odletového i příletového letiště.")
        else:
            map_event = render_folium_navigable(make_route_overview_map(filtered, bool(dark_mode)), height=680, key="route_overview_nav_map_v0355")
            handle_route_map_interaction(map_event)
            render_map_selection(filtered, rates, dark_mode)


def page_rates(rates: pd.DataFrame):
    st.markdown("## Ceník")
    if rates.empty:
        rates = pd.DataFrame(columns=["id", "registration", "aircraft_type", "valid_from", "price_per_hour", "dry_price_per_hour", "source"])
    display = rates.rename(columns={"id":"ID","registration":"Imatrikulace","aircraft_type":"Typ","valid_from":"Od data","price_per_hour":"Cena Kč/h","dry_price_per_hour":"Suchá hodina Kč/h","source":"Zdroj"})
    edited = st.data_editor(display, hide_index=True, use_container_width=True, num_rows="dynamic", disabled=["ID"] if is_admin() else display.columns.tolist(), height=640, column_config={"Cena Kč/h": st.column_config.NumberColumn(format="%.0f Kč"), "Suchá hodina Kč/h": st.column_config.NumberColumn(format="%.0f Kč")})
    if st.button("Uložit ceník", type="primary", disabled=not is_admin()):
        if require_admin():
            save_rates_editor(edited)
            st.success("Ceník uložen."); st.rerun()



def save_rates_editor(edited: pd.DataFrame) -> None:
    with connect() as con:
        con.execute("DELETE FROM rates")
        for _, row in edited.iterrows():
            reg = normalize_text(row.get("Imatrikulace"))
            if not reg:
                continue
            con.execute(
                "INSERT INTO rates (registration, aircraft_type, valid_from, price_per_hour, dry_price_per_hour, source) VALUES (?, ?, ?, ?, ?, ?)",
                (reg.upper(), normalize_text(row.get("Typ")), normalize_text(row.get("Od data")), float(row.get("Cena Kč/h") or 0), float(row.get("Suchá hodina Kč/h") or 0), normalize_text(row.get("Zdroj"))),
            )
        record_audit(con, "save_rates", "rates", None, {"rows": len(edited)})
        con.commit()
    auto_backup_after_change("save_rates")


def save_aircraft_editor(edited: pd.DataFrame) -> None:
    with connect() as con:
        con.execute("DELETE FROM aircraft")
        now = _now_iso()
        for _, row in edited.iterrows():
            reg = normalize_text(row.get("Imatrikulace"))
            if not reg:
                continue
            con.execute(
                """
                INSERT INTO aircraft (registration, aircraft_type, icao_type, aircraft_class, evidence, default_price_per_hour, active, note, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (reg.upper(), normalize_text(row.get("Typ")), normalize_text(row.get("ICAO typ")), normalize_text(row.get("Třída")), normalize_text(row.get("Evidence")), float(row.get("Výchozí Kč/h") or 0), int(row.get("Aktivní") or 0), normalize_text(row.get("Poznámka")), now, now),
            )
        record_audit(con, "save_aircraft", "aircraft", None, {"rows": len(edited)})
        con.commit()
    auto_backup_after_change("save_aircraft")


def upsert_airport_form(data: dict[str, Any]) -> None:
    ident = _clean_ident(data.get("ident"))
    if not ident:
        raise ValueError("Ident letiště je povinný.")
    lat = _to_float(data.get("latitude_deg"))
    lon = _to_float(data.get("longitude_deg"))
    if lat is None or lon is None:
        raise ValueError("Latitude a longitude jsou povinné.")
    now = _now_iso()
    with connect() as con:
        con.execute(
            """
            INSERT INTO airports (ident, name, airport_type, iso_country, iso_region, municipality, latitude_deg, longitude_deg, elevation_ft, gps_code, iata_code, local_code, source, active, closed, data_quality, imported_at, updated_at, raw_json)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(ident) DO UPDATE SET
                name=excluded.name,
                airport_type=excluded.airport_type,
                iso_country=excluded.iso_country,
                iso_region=excluded.iso_region,
                municipality=excluded.municipality,
                latitude_deg=excluded.latitude_deg,
                longitude_deg=excluded.longitude_deg,
                elevation_ft=excluded.elevation_ft,
                gps_code=excluded.gps_code,
                iata_code=excluded.iata_code,
                local_code=excluded.local_code,
                source=excluded.source,
                active=excluded.active,
                closed=excluded.closed,
                data_quality=excluded.data_quality,
                updated_at=excluded.updated_at,
                raw_json=excluded.raw_json
            """,
            (
                ident,
                normalize_text(data.get("name")),
                normalize_text(data.get("airport_type")) or "manual_field",
                normalize_text(data.get("iso_country")) or "CZ",
                normalize_text(data.get("iso_region")),
                normalize_text(data.get("municipality")),
                lat,
                lon,
                _to_float(data.get("elevation_ft")),
                normalize_text(data.get("gps_code")),
                normalize_text(data.get("iata_code")),
                normalize_text(data.get("local_code")),
                normalize_text(data.get("source")) or "manual",
                int(bool(data.get("active", True))),
                int(bool(data.get("closed", False))),
                normalize_text(data.get("data_quality")) or "manual",
                now,
                now,
                json.dumps(data, ensure_ascii=False, default=str),
            ),
        )
        record_audit(con, "upsert_airport", "airports", ident, data)
        con.commit()
    auto_backup_after_change("upsert_airport")


def page_database():
    st.markdown("## Databáze")
    airports = read_airports(active_only=False)
    aircraft = read_table("aircraft")
    tracks = read_table("flight_tracks")
    points = read_table("track_points")
    metas = read_table("app_meta")
    audits = read_table("audit_log") if "audit_log" else pd.DataFrame()

    c1, c2, c3, c4 = st.columns(4)
    with c1: metric_card("Letiště / plochy", str(len(airports)), "databázová tabulka")
    with c2: metric_card("Letadla", str(len(aircraft)), "registrace")
    with c3: metric_card("Tracky", str(len(tracks)), "KML soubory")
    with c4: metric_card("GPS body", f"{len(points):,}".replace(",", " "), "normalizováno")

    tab_airports, tab_aircraft, tab_backup, tab_meta = st.tabs(["Letiště", "Letadla", "Záloha", "Meta"])
    with tab_airports:
        col1, col2, col3 = st.columns([1,1,2])
        with col1:
            country = st.selectbox("Země", ["Vše"] + sorted([x for x in airports.get("iso_country", pd.Series(dtype=str)).dropna().unique() if x]), index=0)
        with col2:
            source = st.selectbox("Zdroj", ["Vše"] + sorted([x for x in airports.get("source", pd.Series(dtype=str)).dropna().unique() if x]), index=0)
        with col3:
            q = st.text_input("Hledat", value="")
        view = airports.copy()
        if country != "Vše":
            view = view[view["iso_country"].eq(country)]
        if source != "Vše":
            view = view[view["source"].eq(source)]
        if q.strip():
            ql = q.strip().lower()
            mask = (
                view["ident"].fillna("").str.lower().str.contains(ql)
                | view["name"].fillna("").str.lower().str.contains(ql)
                | view["municipality"].fillna("").str.lower().str.contains(ql)
            )
            view = view[mask]
        cols = ["ident","name","airport_type","iso_country","municipality","latitude_deg","longitude_deg","source","data_quality","active","closed"]
        st.dataframe(view[[c for c in cols if c in view.columns]].head(1000), hide_index=True, use_container_width=True, height=430)
        st.download_button("Export letišť CSV", data=airports.to_csv(index=False).encode("utf-8"), file_name="airports_export.csv", mime="text/csv", use_container_width=True)
        st.markdown("### Přidat / upravit letiště")
        if not is_admin():
            st.info("Ruční editace letišť je dostupná jen pro admina.")
        else:
            with st.form("airport_upsert_form"):
                a1, a2, a3 = st.columns(3)
                with a1:
                    ident = st.text_input("Ident", value="").upper()
                    name = st.text_input("Název", value="")
                    airport_type = st.selectbox("Typ", ["ultralight_field", "small_airport", "medium_airport", "large_airport", "heliport", "closed", "manual_field"], index=0)
                with a2:
                    iso_country = st.text_input("Země", value="CZ").upper()
                    iso_region = st.text_input("Region", value="")
                    municipality = st.text_input("Obec", value="")
                with a3:
                    latitude_deg = st.number_input("Latitude", value=50.0, format="%.6f")
                    longitude_deg = st.number_input("Longitude", value=14.0, format="%.6f")
                    elevation_ft = st.number_input("Elevation ft", value=0.0, format="%.0f")
                active = st.checkbox("Aktivní", value=True)
                closed = st.checkbox("Uzavřené", value=False)
                submitted = st.form_submit_button("Uložit letiště / plochu", type="primary")
            if submitted:
                try:
                    upsert_airport_form({"ident": ident, "name": name, "airport_type": airport_type, "iso_country": iso_country, "iso_region": iso_region, "municipality": municipality, "latitude_deg": latitude_deg, "longitude_deg": longitude_deg, "elevation_ft": elevation_ft, "source": "manual", "active": active, "closed": closed, "data_quality": "manual"})
                    st.success("Letiště uloženo."); st.rerun()
                except Exception as exc:
                    st.error(str(exc))

    with tab_aircraft:
        if aircraft.empty:
            aircraft = pd.DataFrame(columns=["id","registration","aircraft_type","icao_type","aircraft_class","evidence","default_price_per_hour","active","note"])
        display = aircraft.rename(columns={"id":"ID","registration":"Imatrikulace","aircraft_type":"Typ","icao_type":"ICAO typ","aircraft_class":"Třída","evidence":"Evidence","default_price_per_hour":"Výchozí Kč/h","active":"Aktivní","note":"Poznámka"})
        cols = ["ID","Imatrikulace","Typ","ICAO typ","Třída","Evidence","Výchozí Kč/h","Aktivní","Poznámka"]
        display = display[[c for c in cols if c in display.columns]]
        edited = st.data_editor(display.sort_values("Imatrikulace") if not display.empty else display, hide_index=True, use_container_width=True, num_rows="dynamic", disabled=["ID"] if is_admin() else display.columns.tolist(), height=560, column_config={"Výchozí Kč/h": st.column_config.NumberColumn(format="%.0f Kč"), "Aktivní": st.column_config.CheckboxColumn()})
        if st.button("Uložit letadla", type="primary", disabled=not is_admin()):
            if require_admin():
                save_aircraft_editor(edited)
                st.success("Letadla uložena."); st.rerun()

    # Import světové databáze byl odstraněn z běžného UI po prvotním naplnění tabulky airports.
    # Importní funkce zůstávají v kódu pro případ budoucí servisní migrace, ale nejsou vystavené v aplikaci.

    with tab_backup:
        st.markdown("### SQLite + GitHub backup")
        dirty = ""
        last_change = ""
        last_backup = ""
        if not metas.empty:
            md = dict(zip(metas["key"], metas["value"]))
            dirty = md.get("dirty", "")
            last_change = md.get("last_change_at", "")
            last_backup = md.get("last_github_backup_at", "")
        b1, b2, b3 = st.columns(3)
        with b1: metric_card("Stav", "Nezálohováno" if dirty == "1" else "OK", "dirty flag")
        with b2: metric_card("Poslední změna", last_change[:19] if last_change else "—", "UTC")
        with b3: metric_card("GitHub backup", last_backup[:19] if last_backup else "—", "UTC")
        if github_auto_backup_enabled():
            st.success("Automatická GitHub záloha je zapnutá. Po každé potvrzené změně se databáze uloží do repozitáře.")
        elif github_backup_configured():
            st.warning("GitHub token je nastavený, ale automatická záloha je vypnutá. Zapni github.auto_backup = true v Secrets.")
        else:
            st.warning("Automatická GitHub záloha není nastavená. Změny ve Streamlit Cloud mohou po restartu zmizet.")
        if st.session_state.get("last_auto_backup_status") == "ok":
            pass
        elif st.session_state.get("last_auto_backup_status") == "error":
            st.error(f"Poslední automatická záloha selhala: {st.session_state.get('last_auto_backup_error')}")
        with open(DB_PATH, "rb") as f:
            st.download_button("Stáhnout SQLite databázi", f.read(), file_name="logbook.sqlite", use_container_width=True)
        if github_backup_configured():
            if st.button("Uložit aktuální databázi na GitHub", type="primary", disabled=not is_admin(), use_container_width=True):
                if require_admin():
                    try:
                        url = backup_database_to_github()
                        st.success("Databáze zazálohována na GitHub." + (f" Commit: {url}" if url else ""))
                    except Exception as exc:
                        st.error(f"Backup selhal: {exc}")
        else:
            st.info("GitHub backup není nakonfigurovaný ve Streamlit Secrets. Stále můžeš ručně stahovat SQLite soubor.")
        restore = st.file_uploader("Obnovit SQLite databázi ze souboru", type=["sqlite", "db"], key="restore_db_upload")
        confirm = st.text_input("Pro obnovení napiš OBNOVIT", value="")
        if restore is not None and st.button("Obnovit databázi", disabled=not is_admin() or confirm != "OBNOVIT", use_container_width=True):
            if require_admin():
                try:
                    restore_database_from_upload(restore)
                    st.success("Databáze obnovena."); st.rerun()
                except Exception as exc:
                    st.error(f"Obnova selhala: {exc}")

    with tab_meta:
        st.markdown("### Metadata")
        if metas.empty:
            st.info("Žádná metadata.")
        else:
            st.dataframe(metas.sort_values("key"), hide_index=True, use_container_width=True)
        st.markdown("### Audit log")
        if audits.empty:
            st.info("Žádný audit log.")
        else:
            st.dataframe(audits.sort_values("id", ascending=False).head(500), hide_index=True, use_container_width=True, height=360)


def make_control_df(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, r in df.iterrows():
        issues = []
        if pd.isna(r.get("date_dt")): issues.append("datum")
        else:
            y = int(r["date_dt"].year)
            if y < 2000 or y > 2035: issues.append("rok")
        for col, label in [("evidence","evidence"),("registration","imatrikulace"),("role","funkce")]:
            if not str(r.get(col) or "").strip(): issues.append(label)
        if r.get("starts", 0) <= 0: issues.append("starty")
        if r.get("block_minutes") is None or pd.isna(r.get("block_minutes")): issues.append("block time")
        if r.get("air_minutes") is None or pd.isna(r.get("air_minutes")): issues.append("air time")
        if pd.notna(r.get("air_minutes")) and pd.notna(r.get("block_minutes")) and r["air_minutes"] > r["block_minutes"]: issues.append("air > block")
        if pd.isna(r.get("price_per_hour")) or float(r.get("price_per_hour") or 0) <= 0: issues.append("sazba")
        if issues:
            rows.append({"ID": r.get("id"), "Datum": r.get("date"), "Imatrikulace": r.get("registration"), "Problém": ", ".join(issues)})
    return pd.DataFrame(rows)


def page_control(df: pd.DataFrame):
    st.markdown("## Kontrola")
    control = make_control_df(df)
    if control.empty:
        st.success("OK")
    else:
        st.dataframe(control, hide_index=True, use_container_width=True)


def export_excel(df: pd.DataFrame) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Zápisník letů"
    headers = ["Datum","Evidence","Imatrikulace","Typ","Třída","Odlet","Přílet","Off Block","Takeoff","Landing","On Block","Block Time","Air Time","Starty","Velitel","Instruktor","Funkce","Úloha","Cena Kč/h","Cena letu Kč","GPS tracky","GPS km","Poznámka"]
    ws.append(headers)
    work = df.sort_values(["date_dt","id"]).copy()
    for _, r in work.iterrows():
        ws.append([r.get("date"), r.get("evidence"), r.get("registration"), r.get("aircraft_type"), r.get("aircraft_class"), r.get("departure"), r.get("arrival"), r.get("off_block"), r.get("takeoff"), r.get("landing"), r.get("on_block"), fmt_minutes(r.get("block_minutes")), fmt_minutes(r.get("air_minutes")), int(r.get("starts") or 0), r.get("commander"), r.get("instructor"), r.get("role"), r.get("task"), float(r.get("price_per_hour") or 0), float(r.get("cost") or 0), int(r.get("track_count") or 0), float(r.get("gps_km") or 0), r.get("note")])
    style_worksheet(ws)
    ws2 = wb.create_sheet("Souhrny")
    summary = build_summary(df)
    rows = [["Metrika","Hodnota"],["Celkový nálet",fmt_minutes(summary["total"])],["Air Time",fmt_minutes(summary["air"])],["PIC celkem",fmt_minutes(summary["pic"])],["PIC ULL",fmt_minutes(summary["pic_ull"])],["PIC EASA",fmt_minutes(summary["pic_easa"])],["DUAL",fmt_minutes(summary["dual"])],["Safety Pilot",fmt_minutes(summary["safety"])],["ULL celkem",fmt_minutes(summary["ull"])],["EASA celkem",fmt_minutes(summary["easa"])],["Starty",summary["starts"]],["GPS tracky",summary["tracks"]],["GPS km",summary["gps_km"]],["Náklady",summary["cost"]]]
    for row in rows: ws2.append(row)
    style_worksheet(ws2)
    out = BytesIO(); wb.save(out); return out.getvalue()


def style_worksheet(ws):
    header_fill = PatternFill("solid", fgColor="0F172A"); header_font = Font(color="FFFFFF", bold=True); thin = Side(style="thin", color="D1D5DB")
    for cell in ws[1]:
        cell.fill = header_fill; cell.font = header_font; cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col); max_len = max(len(str(ws.cell(row, col).value or "")) for row in range(1, min(ws.max_row, 200) + 1)); ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 26)


def page_export(df: pd.DataFrame):
    st.markdown("## Export")
    output = export_excel(df)
    c1, c2 = st.columns(2)
    with c1: st.download_button("Stáhnout Excel export", data=output, file_name="export_letovy_zapisnik.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary", use_container_width=True)
    with c2:
        with open(DB_PATH, "rb") as f:
            st.download_button("Stáhnout SQLite databázi", f.read(), file_name="logbook.sqlite", use_container_width=True)

def render_sidebar_toggle() -> None:
    """One smooth sidebar toggle controlled in the browser, without Streamlit rerun."""
    components.html(
        """
        <script>
        (function() {
          const doc = window.parent.document;
          const btnId = 'lb-sidebar-toggle';
          const storageKey = 'lb_sidebar_hidden_v3';

          function hideNativeButtons() {
            const selectors = [
              '[data-testid="stSidebarHeader"]',
              '[data-testid="stSidebarCollapseButton"]',
              '[data-testid="stSidebarCollapsedControl"]',
              '[data-testid="collapsedControl"]',
              'section[data-testid="stSidebar"] button[kind="headerNoPadding"]',
              'section[data-testid="stSidebar"] button[data-testid="baseButton-headerNoPadding"]',
              'button[title*="sidebar" i]',
              'button[aria-label*="sidebar" i]',
              'button[title*="Collapse" i]',
              'button[aria-label*="Collapse" i]',
              'button[title*="Close" i]',
              'button[aria-label*="Close" i]'
            ];
            selectors.forEach(sel => {
              try {
                doc.querySelectorAll(sel).forEach(el => {
                  if (el.id !== btnId) {
                    el.style.setProperty('display', 'none', 'important');
                    el.style.setProperty('visibility', 'hidden', 'important');
                    el.style.setProperty('pointer-events', 'none', 'important');
                  }
                });
              } catch(e) {}
            });
          }

          function ensureButton() {
            let btn = doc.getElementById(btnId);
            if (!btn) {
              btn = doc.createElement('button');
              btn.id = btnId;
              btn.type = 'button';
              btn.setAttribute('aria-label', 'Skrýt nebo zobrazit menu');
              btn.title = 'Skrýt / zobrazit menu';
              doc.body.appendChild(btn);
              btn.addEventListener('click', function(ev) {
                ev.preventDefault();
                setHidden(!doc.body.classList.contains('lb-sidebar-hidden'));
              });
            }
            return btn;
          }

          function setHidden(hidden) {
            doc.body.classList.toggle('lb-sidebar-hidden', hidden);
            try { window.parent.localStorage.setItem(storageKey, hidden ? '1' : '0'); } catch(e) {}
            const btn = ensureButton();
            btn.textContent = hidden ? '›' : '‹';
          }

          hideNativeButtons();
          const saved = (function() {
            try { return window.parent.localStorage.getItem(storageKey) === '1'; }
            catch(e) { return false; }
          })();
          setHidden(saved);
          if (!window.parent.__lbSidebarObserver) {
            window.parent.__lbSidebarObserver = new MutationObserver(hideNativeButtons);
            window.parent.__lbSidebarObserver.observe(doc.body, {childList: true, subtree: true});
          }
          setTimeout(hideNativeButtons, 250);
          setTimeout(hideNativeButtons, 1000);
        })();
        </script>
        """,
        height=0,
        width=0,
    )


def render_page_transition_runtime() -> None:
    """Install a tiny front-end page loader.

    Streamlit reruns the Python script after every sidebar button click. Without a
    front-end transition, the previous page visually disappears piece by piece
    while the new page is being generated. This overlay hides that intermediate
    state and makes navigation feel much closer to a normal web app.
    """
    components.html(
        """
        <script>
        (function() {
          const doc = window.parent.document;
          const overlayId = 'lb-page-loader';

          function ensureOverlay() {
            let el = doc.getElementById(overlayId);
            if (!el) {
              el = doc.createElement('div');
              el.id = overlayId;
              el.innerHTML = '<div class="lb-plane-spinner" title="Načítám"><span>✈</span></div>';
              doc.body.appendChild(el);
            }
            return el;
          }
          function showLoader() {
            ensureOverlay();
            doc.body.classList.add('lb-page-loading');
          }
          function hideLoader() {
            ensureOverlay();
            doc.body.classList.remove('lb-page-loading');
          }

          if (!window.parent.__lbPageLoaderInstalled) {
            doc.addEventListener('click', function(ev) {
              const target = ev.target;
              if (!target) return;
              const btn = target.closest && target.closest('button');
              const link = target.closest && target.closest('a');
              if (btn && btn.id !== 'lb-sidebar-toggle') {
                const inSidebar = btn.closest('section[data-testid="stSidebar"]');
                const txt = (btn.innerText || btn.textContent || '').trim();
                if (inSidebar && txt && txt !== 'Odhlásit') showLoader();
              }
              if (link && link.href && link.href.indexOf('flight_id=') !== -1) {
                showLoader();
              }
            }, true);
            window.parent.__lbPageLoaderInstalled = true;
          }
          // The new page has reached the browser once this component runs.
          setTimeout(hideLoader, 120);
          setTimeout(hideLoader, 800);
        })();
        </script>
        """,
        height=0,
        width=0,
    )


def render_page_loaded_signal() -> None:
    """Hide the front-end loader after the current Streamlit page has rendered."""
    components.html(
        """
        <script>
        (function() {
          const doc = window.parent.document;
          function hideLoader() { doc.body.classList.remove('lb-page-loading'); }
          setTimeout(hideLoader, 80);
          setTimeout(hideLoader, 450);
        })();
        </script>
        """,
        height=0,
        width=0,
    )


# -----------------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------------

def main():
    st.set_page_config(page_title="Letový zápisník", layout="wide", initial_sidebar_state="expanded")
    with connect(): pass
    if "page" not in st.session_state:
        st.session_state["page"] = "Dashboard"
    q_flight_id = _query_param_value("flight_id")
    if q_flight_id and str(q_flight_id).isdigit():
        qid = int(q_flight_id)
        if st.session_state.get("dismissed_flight_id") != qid:
            st.session_state["page"] = "Lety"
            st.session_state["open_flight_dialog_id"] = qid
            st.session_state["selected_flight_id"] = qid
    q_map_airport = _query_param_value("map_airport")
    q_map_route = _query_param_value("map_route")
    if q_map_airport:
        st.session_state["page"] = "Mapa"
        st.session_state["map_airport"] = str(q_map_airport).upper().strip()
        st.session_state["map_mode_v035"] = "Orientační mapa letišť"
        st.session_state.pop("map_route", None)
    elif q_map_route:
        st.session_state["page"] = "Mapa"
        st.session_state["map_route"] = str(q_map_route).upper().strip()
        st.session_state["map_mode_v035"] = "Orientační mapa letišť"
        st.session_state.pop("map_airport", None)
    with st.sidebar:
        # Aplikace běží trvale v tmavém režimu; přepínač je z finálního UI odstraněn.
        dark_mode = True
        st.markdown("## Letový zápisník")
        st.markdown(f'<div class="sidebar-version">{APP_VERSION}</div>', unsafe_allow_html=True)
        render_sidebar_nav()
        render_auth_sidebar()
    apply_ui_theme(dark_mode)
    render_sidebar_toggle()
    render_page_transition_runtime()
    app_header()
    page = st.session_state.get("page", "Dashboard")
    flights = read_flights()
    rates = read_table("rates")
    if not rates.empty:
        rates["registration"] = rates["registration"].fillna("").str.upper()
    if page == "Dashboard": page_dashboard(flights)
    elif page == "Lety": page_logbook(flights, rates, dark_mode)
    elif page == "Nový let": page_new_flight(rates, dark_mode)
    elif page == "Mapa": page_maps(flights, rates, dark_mode)
    elif page == "Ceník": page_rates(rates)
    elif page == "Databáze": page_database()
    elif page == "Kontrola":
        # Legacy route: stránka kontroly už není v navigaci, ale starý stav relace může existovat.
        st.session_state["page"] = "Dashboard"
        st.rerun()
    elif page == "Export": page_export(flights)
    render_page_loaded_signal()

if __name__ == "__main__":
    main()
